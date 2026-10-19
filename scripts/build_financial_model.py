#!/usr/bin/env python3
"""Generate study/feasibility/financial-model.xlsx, the farm's financial model.

    uv run --with openpyxl python scripts/build_financial_model.py

Every input is in the tables below and lands in blue on the Control, Assumptions,
Models, Curves and Sensitivity sheets. Every other cell is a live formula, so the
workbook recalculates in Excel when any input changes.

Evidence labels follow CLAUDE.md: V verified, E estimate, A assumption, U unknown.
Inputs labelled A or U are placeholders until research or field data replaces them.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "study" / "feasibility" / "financial-model.xlsx"

MONTHS = 120  # the monthly engine covers 10 farm years
YEARS = 10
MC0 = 4  # monthly sheets: column D is month 1
YC0 = 4  # annual sheets: column D is year 1

# Structural constants. They shape the formulas (fixed lags), so change them here and rebuild.
SALE_AGE_M = 9  # months from birth to sale of young stock
FIRST_BIRTH_LAG_M = 6  # months between a female joining the herd and her first birth
RECEIVABLE_LAG_M = 1  # months between a credit sale and its collection
DAYS_PER_MONTH = 30.42

# ---------------------------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------------------------

PENDING = "placeholder — pending Stage 1 research"

# Control sheet: (key, label, value, unit, note)
CONTROL = [
    ("scenario", "Scenario (1 downside · 2 base · 3 upside)", 2, "index", "Drives every Assumptions row"),
    ("model", "Farm model (1 M1 · 2 M2 · 3 M3)", 1, "index", "Drives every Models row"),
    ("area", "Total farm area", 5, "ha", "[A] founder: 2–10 ha; base 5 ha until surveyed"),
    ("capital", "Capital available, years 1–2", 35_000_000, "RWF", "[A] founder: under USD 25k"),
    ("start_year", "Model start year", 2026, "year", "Month 1 = first month of spending"),
    ("start_month", "Model start month (1–12)", 11, "month", "11 = November"),
    ("plant_m", "Tree planting month (model month #)", 5, "month #", "[A] 5 = March 2027, Season B onset"),
    ("past_m", "Pasture/fodder planting month (model month #)", 5, "month #", "[A] same rains as trees"),
]

# Assumptions sheet: (id, group, parameter, unit, downside, base, upside, label, source, note)
# Downside holds the pessimistic value whichever direction that is (a higher cost, a lower price).
A = [
    # General
    ("FX", "General", "Exchange rate", "RWF per USD", 1470, 1470, 1470, "A", PENDING, "Verify against the BNR reference rate"),
    ("DR", "General", "Real discount rate", "per year", 0.14, 0.12, 0.10, "A", "judgement", "Real terms: the model is in constant 2026 RWF"),
    ("TAX", "General", "Income tax rate on profit", "share", 0.28, 0.28, 0.28, "U", PENDING, "Confirm agricultural income treatment with RRA"),
    ("CREDIT", "General", "Share of sales on credit", "share", 0.6, 0.5, 0.3, "A", "judgement", f"Collected {RECEIVABLE_LAG_M} month later"),
    ("CONT", "General", "Capex contingency", "share of capex", 0.15, 0.10, 0.08, "A", "judgement", ""),
    ("MKTG", "General", "Marketing and sales cost", "share of sales", 0.03, 0.02, 0.02, "A", "judgement", ""),
    ("RM", "General", "Repairs and maintenance", "share of capex per yr", 0.05, 0.04, 0.03, "A", "judgement", ""),
    # Labour
    ("WAGE", "Labour", "Permanent worker wage", "RWF per month", 60_000, 50_000, 45_000, "A", PENDING, ""),
    ("DAYWAGE", "Labour", "Casual labour wage", "RWF per day", 2_500, 2_000, 1_800, "A", PENDING, ""),
    ("DAYS_HA", "Labour", "Casual labour in orchard", "days per ha per yr", 120, 100, 80, "A", PENDING, ""),
    # Capital costs
    ("FENCE", "Capital costs", "Perimeter fence", "RWF per m", 6_000, 5_000, 4_000, "A", PENDING, ""),
    ("IFENCE", "Capital costs", "Internal paddock fence", "RWF per m", 3_500, 3_000, 2_500, "A", PENDING, ""),
    ("GATE", "Capital costs", "Gate", "RWF each", 250_000, 200_000, 150_000, "A", PENDING, ""),
    ("TANK", "Capital costs", "Water storage", "RWF per m3", 120_000, 100_000, 90_000, "A", PENDING, ""),
    ("GUTTER", "Capital costs", "Rainwater harvesting fittings", "RWF lump", 400_000, 300_000, 250_000, "A", PENDING, ""),
    ("BORE", "Capital costs", "Borehole (if legal and viable)", "RWF lump", 12_000_000, 9_000_000, 7_000_000, "U", PENDING, "Hydrogeology and permit unknown"),
    ("PUMP", "Capital costs", "Solar pump system", "RWF lump", 4_500_000, 3_500_000, 3_000_000, "A", PENDING, ""),
    ("DRIP", "Capital costs", "Drip irrigation", "RWF per ha", 4_500_000, 3_500_000, 3_000_000, "A", PENDING, ""),
    ("PAST_EST", "Capital costs", "Pasture establishment", "RWF per ha", 600_000, 450_000, 350_000, "A", PENDING, ""),
    ("FOD_EST", "Capital costs", "Fodder bank establishment", "RWF per ha", 900_000, 700_000, 550_000, "A", PENDING, ""),
    ("HOUSE", "Capital costs", "Animal house", "RWF per m2", 60_000, 45_000, 35_000, "A", PENDING, ""),
    ("M2HEAD", "Capital costs", "House floor area per adult animal", "m2 per head", 1.5, 1.2, 1.0, "A", PENDING, ""),
    ("QUAR", "Capital costs", "Quarantine pen and handling race", "RWF lump", 800_000, 600_000, 500_000, "A", PENDING, ""),
    ("SOIL", "Capital costs", "Soil test", "RWF per sample", 40_000, 30_000, 25_000, "U", PENDING, "Get lab quotes"),
    ("WATER", "Capital costs", "Water test", "RWF per sample", 60_000, 45_000, 35_000, "U", PENDING, "Get lab quotes"),
    ("SURVEY", "Capital costs", "Survey, title and legal checks", "RWF lump", 600_000, 400_000, 300_000, "U", PENDING, ""),
    ("PERMIT", "Capital costs", "Permits and compliance", "RWF lump", 500_000, 300_000, 200_000, "U", PENDING, ""),
    ("PROF", "Capital costs", "Professional set-up visits (agronomist, vet)", "RWF lump", 1_000_000, 800_000, 600_000, "A", PENDING, ""),
    ("PV", "Capital costs", "Solar PV with battery", "RWF per kWp", 2_500_000, 2_000_000, 1_700_000, "A", PENDING, ""),
    ("SCALE", "Capital costs", "Livestock weighing scale", "RWF lump", 600_000, 450_000, 350_000, "A", PENDING, ""),
    # Operating costs
    ("VET", "Operating costs", "Veterinary care and medicines", "RWF per head per yr", 8_000, 6_000, 5_000, "A", PENDING, "Confirm with a Rwandan vet"),
    ("MINERAL", "Operating costs", "Minerals and supplements", "RWF per head per yr", 3_000, 2_500, 2_000, "A", PENDING, ""),
    ("FERT", "Operating costs", "Orchard inputs (fertiliser, compost, protection)", "RWF per ha per yr", 900_000, 700_000, 550_000, "A", PENDING, ""),
    ("PASTM", "Operating costs", "Pasture and fodder upkeep", "RWF per ha per yr", 150_000, 100_000, 80_000, "A", PENDING, ""),
    ("TRANS_KG", "Operating costs", "Transport of fruit to buyers", "RWF per kg", 80, 60, 40, "U", PENDING, "Depends on the unconfirmed location"),
    ("TRANS_HD", "Operating costs", "Transport of livestock to buyers", "RWF per head", 3_000, 2_000, 1_500, "U", PENDING, "Depends on the unconfirmed location"),
    ("PACK", "Operating costs", "Packaging", "RWF per kg fruit", 40, 30, 25, "A", PENDING, ""),
    ("COMMS", "Operating costs", "Communications and software", "RWF per month", 40_000, 30_000, 25_000, "A", PENDING, ""),
    ("UTIL", "Operating costs", "Utilities and fuel", "RWF per month", 40_000, 30_000, 20_000, "A", PENDING, ""),
    ("ADMIN", "Operating costs", "Administration and accounting", "RWF per month", 60_000, 40_000, 30_000, "A", PENDING, ""),
    ("INS", "Operating costs", "Livestock insurance premium", "share of herd value per yr", 0.06, 0.05, 0.04, "U", PENDING, "Check the national agriculture insurance scheme"),
    ("LANDFEE", "Operating costs", "Land fees and taxes", "RWF per yr", 200_000, 120_000, 80_000, "U", PENDING, "Depends on the title"),
    # Feed
    ("PAST_Y", "Feed", "Pasture yield", "t DM per ha per yr", 6, 9, 12, "A", PENDING, ""),
    ("FOD_Y", "Feed", "Fodder bank yield (cut-and-carry)", "t DM per ha per yr", 12, 18, 24, "A", PENDING, ""),
    ("GRAZE", "Feed", "Grazing utilisation", "share", 0.50, 0.60, 0.65, "A", PENDING, ""),
    ("CUT", "Feed", "Cut-and-carry utilisation", "share", 0.70, 0.80, 0.85, "A", PENDING, ""),
    ("CONS_EFF", "Feed", "Hay/silage conservation efficiency", "share", 0.60, 0.70, 0.80, "A", PENDING, ""),
    ("CONS_COST", "Feed", "Conservation cost", "RWF per kg DM", 30, 20, 15, "A", PENDING, ""),
    ("FEED_COST", "Feed", "Purchased feed", "RWF per kg DM", 300, 220, 180, "A", PENDING, ""),
    ("READY", "Feed", "Months from planting to first forage", "months", 6, 5, 4, "A", PENDING, ""),
]

# Livestock parameters per species: (suffix, parameter, unit, goats d/b/u, sheep d/b/u, label, note)
LIVESTOCK = [
    ("PART", "Births (parturitions) per breeding female", "per yr", (1.1, 1.3, 1.5), (0.9, 1.1, 1.3), "A", ""),
    ("LIT", "Litter size", "young per birth", (1.3, 1.5, 1.7), (1.1, 1.2, 1.4), "A", ""),
    ("YMORT", "Young mortality to weaning", "share", (0.25, 0.15, 0.10), (0.25, 0.15, 0.10), "A", ""),
    ("AMORT", "Adult mortality", "share per yr", (0.10, 0.06, 0.04), (0.10, 0.06, 0.04), "A", ""),
    ("CULL", "Adult cull rate", "share per yr", (0.15, 0.15, 0.15), (0.15, 0.15, 0.15), "A", ""),
    ("SALEWT", f"Live weight at sale ({SALE_AGE_M} months)", "kg", (20, 24, 28), (22, 26, 30), "A", ""),
    ("ADULTWT", "Adult live weight (culls)", "kg", (32, 35, 38), (30, 33, 36), "A", ""),
    ("PRICE", "Sale price, live weight", "RWF per kg", (3_000, 3_500, 4_000), (2_800, 3_200, 3_600), "A", ""),
    ("FPRICE", "Breeding female purchase price", "RWF per head", (60_000, 50_000, 45_000), (55_000, 45_000, 40_000), "A", ""),
    ("MPRICE", "Breeding male purchase price", "RWF per head", (250_000, 200_000, 150_000), (200_000, 150_000, 120_000), "A", ""),
    ("BPREM", "Price premium for breeding-stock sales", "share", (0.2, 0.3, 0.4), (0.2, 0.3, 0.4), "A", ""),
    ("DMI_A", "Dry-matter intake, adult", "kg DM per day", (1.2, 1.1, 1.0), (1.3, 1.2, 1.1), "A", ""),
    ("DMI_Y", "Dry-matter intake, growing young", "kg DM per day", (0.7, 0.6, 0.55), (0.7, 0.6, 0.55), "A", ""),
    ("RATIO", "Breeding females per male", "females", (25, 25, 25), (30, 30, 30), "A", ""),
    ("FSHARE", "Female share of births", "share", (0.5, 0.5, 0.5), (0.5, 0.5, 0.5), "A", ""),
]
SPECIES = [("GT", "Goats"), ("SH", "Sheep")]

# Crops: (key, name, density/ha, survival d/b/u, seedling RWF d/b/u, planting RWF d/b/u,
#         yield multiplier d/b/u, marketable share d/b/u, farm-gate RWF/kg d/b/u, replant cycle yrs,
#         yield curve kg per plant Y1..Y10)
# The yield curve assumes planting in year 1 at the Season B rains; cycle 0 = perennial.
CROPS = [
    ("AVO", "Avocado (grafted)", 204, (0.80, 0.88, 0.93), (3_000, 2_500, 2_000), (1_500, 1_200, 1_000),
     (0.7, 1.0, 1.2), (0.75, 0.85, 0.90), (300, 400, 500), 0, [0, 0, 5, 15, 30, 50, 70, 90, 100, 110]),
    ("MAN", "Mango (grafted)", 156, (0.80, 0.88, 0.93), (2_500, 2_000, 1_500), (1_500, 1_200, 1_000),
     (0.7, 1.0, 1.2), (0.70, 0.80, 0.85), (250, 350, 450), 0, [0, 0, 5, 15, 30, 45, 60, 75, 90, 100]),
    ("PAS", "Passion fruit", 1_111, (0.75, 0.85, 0.90), (500, 300, 200), (2_500, 2_000, 1_600),
     (0.7, 1.0, 1.2), (0.75, 0.85, 0.90), (500, 700, 900), 3, [2, 10, 6, 0, 0, 0, 0, 0, 0, 0]),
    ("PAP", "Papaya", 1_667, (0.75, 0.85, 0.90), (400, 300, 200), (800, 600, 500),
     (0.7, 1.0, 1.2), (0.70, 0.80, 0.85), (200, 300, 400), 3, [5, 30, 20, 0, 0, 0, 0, 0, 0, 0]),
    ("TTO", "Tree tomato", 1_667, (0.75, 0.85, 0.90), (300, 200, 150), (600, 500, 400),
     (0.7, 1.0, 1.2), (0.70, 0.80, 0.85), (300, 400, 500), 4, [0, 10, 15, 10, 0, 0, 0, 0, 0, 0]),
    ("CIT", "Citrus", 400, (0.80, 0.88, 0.93), (2_500, 2_000, 1_500), (1_200, 1_000, 800),
     (0.7, 1.0, 1.2), (0.70, 0.80, 0.85), (250, 350, 450), 0, [0, 0, 5, 15, 30, 45, 60, 70, 80, 90]),
    ("MAC", "Macadamia (nut in shell)", 204, (0.80, 0.88, 0.93), (4_000, 3_500, 3_000), (1_500, 1_200, 1_000),
     (0.7, 1.0, 1.2), (0.85, 0.90, 0.95), (800, 1_000, 1_200), 0, [0, 0, 0, 1, 3, 6, 9, 12, 15, 18]),
]

# Boundary timber: (id, parameter, unit, d/b/u, label, note)
TIMBER = [
    ("T_SHAPE", "Boundary shape factor (1 = square)", "factor", (1.25, 1.15, 1.05), "U", "Site-specific"),
    ("T_SPACING", "Boundary tree spacing", "m", (4, 4, 4), "A", ""),
    ("T_SURV", "Boundary tree survival", "share", (0.75, 0.85, 0.90), "A", ""),
    ("T_SEED", "Boundary tree seedling", "RWF each", (300, 200, 150), "A", ""),
    ("T_PLANT", "Boundary tree planting", "RWF each", (500, 400, 300), "A", ""),
    ("T_POLEYEAR", "Year of first pole thinning", "farm year", (8, 7, 6), "A", ""),
    ("T_POLESHARE", "Share of trees thinned as poles", "share", (0.3, 0.3, 0.3), "A", ""),
    ("T_POLEPRICE", "Pole price", "RWF each", (6_000, 8_000, 10_000), "A", ""),
    ("T_STAND", "Standing value per remaining tree at Y10 (memo)", "RWF each", (25_000, 35_000, 45_000), "A", "Not in NPV"),
]

# Models sheet: (id, group, parameter, unit, M1, M2, M3, label, note)
M = [
    ("Z01", "Land use (share of total area)", "Z0+Z1 base, store, nursery", "share", 0.05, 0.05, 0.07, "A", "See zoning-concept.md"),
    ("Z2", "Land use (share of total area)", "Z2 orchard", "share", 0.25, 0.35, 0.38, "A", ""),
    ("Z3", "Land use (share of total area)", "Z3 paddocks and fodder bank", "share", 0.50, 0.40, 0.35, "A", ""),
    ("Z4", "Land use (share of total area)", "Z4 boundary and shelter trees", "share", 0.10, 0.10, 0.10, "A", ""),
    ("Z5", "Land use (share of total area)", "Z5 lake/wetland buffer, steep, conservation", "share", 0.10, 0.10, 0.10, "U", "Depends on shoreline and survey"),
    ("FODDER", "Land use (share of total area)", "Fodder bank share of Z3", "share", 0.30, 0.30, 0.35, "A", ""),
    ("OR_AVO", "Orchard mix (share of Z2)", "Avocado", "share", 1.00, 0.50, 0.40, "A", "Pending site suitability"),
    ("OR_MAN", "Orchard mix (share of Z2)", "Mango", "share", 0.00, 0.00, 0.10, "A", ""),
    ("OR_PAS", "Orchard mix (share of Z2)", "Passion fruit", "share", 0.00, 0.30, 0.20, "A", ""),
    ("OR_PAP", "Orchard mix (share of Z2)", "Papaya", "share", 0.00, 0.20, 0.10, "A", ""),
    ("OR_TTO", "Orchard mix (share of Z2)", "Tree tomato", "share", 0.00, 0.00, 0.10, "A", ""),
    ("OR_CIT", "Orchard mix (share of Z2)", "Citrus", "share", 0.00, 0.00, 0.10, "A", ""),
    ("OR_MAC", "Orchard mix (share of Z2)", "Macadamia", "share", 0.00, 0.00, 0.00, "A", ""),
    ("GT_T1M", "Livestock introduction", "Goats tranche 1: month", "month #", 9, 9, 9, "A", "After fencing, water, house and forage"),
    ("GT_T1F", "Livestock introduction", "Goats tranche 1: breeding females", "head", 15, 20, 25, "A", ""),
    ("GT_T1MALE", "Livestock introduction", "Goats tranche 1: breeding males", "head", 1, 1, 1, "A", ""),
    ("GT_T2M", "Livestock introduction", "Goats tranche 2: month", "month #", 18, 18, 18, "A", "After quarantine and first results"),
    ("GT_T2F", "Livestock introduction", "Goats tranche 2: breeding females", "head", 10, 15, 20, "A", ""),
    ("GT_T2MALE", "Livestock introduction", "Goats tranche 2: breeding males", "head", 0, 1, 1, "A", ""),
    ("GT_TARGET", "Livestock introduction", "Goats: target breeding females", "head", 40, 50, 60, "A", "Check against the feed balance"),
    ("SH_T1M", "Livestock introduction", "Sheep tranche 1: month", "month #", 0, 12, 12, "A", "0 = no sheep"),
    ("SH_T1F", "Livestock introduction", "Sheep tranche 1: breeding females", "head", 0, 10, 15, "A", ""),
    ("SH_T1MALE", "Livestock introduction", "Sheep tranche 1: breeding males", "head", 0, 1, 1, "A", ""),
    ("SH_T2M", "Livestock introduction", "Sheep tranche 2: month", "month #", 0, 0, 24, "A", "0 = none"),
    ("SH_T2F", "Livestock introduction", "Sheep tranche 2: breeding females", "head", 0, 0, 10, "A", ""),
    ("SH_T2MALE", "Livestock introduction", "Sheep tranche 2: breeding males", "head", 0, 0, 0, "A", ""),
    ("SH_TARGET", "Livestock introduction", "Sheep: target breeding females", "head", 0, 20, 30, "A", ""),
    ("RETAIN", "Livestock management", "Share of female young kept as replacements (while below target)", "share", 0.60, 0.50, 0.50, "A", ""),
    ("BSHARE", "Livestock management", "Share of young stock sold as breeding stock", "share", 0.00, 0.10, 0.20, "A", "Needs records and genetics"),
    ("CONSERVE", "Livestock management", "Share of surplus forage conserved (hay/silage)", "share", 0.30, 0.50, 0.70, "A", ""),
    ("WORKERS", "Operations", "Permanent workers", "people", 1, 2, 3, "A", ""),
    ("SOIL_N", "Capital items", "Soil samples", "samples", 6, 8, 10, "A", "See soil-water-testing-plan.md"),
    ("WATER_N", "Capital items", "Water samples", "samples", 2, 3, 4, "A", ""),
    ("IFENCE_M", "Capital items", "Internal paddock fence", "m per ha of Z3", 250, 300, 350, "A", ""),
    ("GATES", "Capital items", "Gates", "count", 2, 3, 4, "A", ""),
    ("STORAGE", "Capital items", "Water storage", "m3", 10, 20, 40, "A", ""),
    ("BOREHOLE", "Capital items", "Borehole (1 = yes)", "flag", 0, 0, 1, "U", "Only if legal and viable"),
    ("SOLARPUMP", "Capital items", "Solar pump (1 = yes)", "flag", 0, 1, 1, "A", ""),
    ("DRIPON", "Capital items", "Drip on orchard (1 = yes)", "flag", 0, 1, 1, "A", ""),
    ("PV_KWP", "Capital items", "Solar PV size", "kWp", 0.3, 0.6, 2.0, "A", "Lights, sensors, comms"),
    ("SENSORS", "Capital items", "Sensors and controllers", "RWF lump", 150_000, 600_000, 3_000_000, "A", "M1/M2 founder-built; M3 industrial-grade"),
    ("CONNECT", "Capital items", "Connectivity (router, gateway)", "RWF lump", 80_000, 150_000, 600_000, "A", ""),
    ("SOFTWARE", "Capital items", "Software and devices", "RWF lump", 0, 0, 1_500_000, "A", "Founder builds; hosting in opex"),
    ("SECURITY", "Capital items", "Security (lights, cameras)", "RWF lump", 200_000, 500_000, 1_500_000, "A", ""),
    ("NURSERY", "Capital items", "Nursery (shade net, tools)", "RWF lump", 0, 300_000, 800_000, "A", ""),
    ("COOLSTORE", "Capital items", "Cool store / packhouse", "RWF lump", 0, 0, 5_000_000, "A", ""),
    ("FEEDSTORE", "Capital items", "Feed store", "RWF lump", 500_000, 800_000, 1_500_000, "A", ""),
    ("TOOLS", "Capital items", "Tools and small equipment", "RWF lump", 600_000, 900_000, 1_500_000, "A", ""),
]

# Sensitivity multipliers: (key, driver, default, how it is applied)
SENS = [
    ("S_PRICE_FRUIT", "Fruit prices", 1.0, "Multiplies every fruit farm-gate price"),
    ("S_PRICE_LIVE", "Livestock prices", 1.0, "Multiplies the live-weight price of goats and sheep"),
    ("S_YIELD", "Fruit yields", 1.0, "Multiplies every yield curve"),
    ("S_MORT", "Livestock mortality", 1.0, "Multiplies young and adult mortality"),
    ("S_FEED", "Purchased feed cost", 1.0, "Multiplies the purchased-feed price"),
    ("S_LABOR", "Labour cost", 1.0, "Multiplies permanent and casual wages"),
    ("S_CAPEX", "Capital costs", 1.0, "Multiplies every capital item, seedling and animal purchase"),
    ("S_TRANS", "Transport cost", 1.0, "Multiplies transport per kg and per head"),
    ("S_DELAY", "Delay to tree maturity", 0, "Years added before perennial trees start their yield curve"),
]

# Monthly calendar inputs (Curves sheet). Seasonality = share of a year's harvest per calendar month.
UNIFORM = [round(1 / 12, 6)] * 11
UNIFORM.append(round(1 - sum(UNIFORM), 6))
SEASON = {c[0]: UNIFORM for c in CROPS}
# Pasture growth by calendar month, averaging 1.0: follows the bimodal rains (Mar–May, Sep–Dec).
PASTURE_GROWTH = [0.9, 0.8, 1.3, 1.5, 1.4, 0.6, 0.4, 0.4, 0.8, 1.2, 1.4, 1.3]
MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# ---------------------------------------------------------------------------------------------
# Styles and helpers
# ---------------------------------------------------------------------------------------------

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_TOTAL = PatternFill("solid", fgColor="F2F2F2")

NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
YEARFMT = '"Y"0'
MONTHFMT = '"M"0'
DATEFMT = "mmm yy"

REF: dict[str, str] = {}  # key -> absolute single-cell reference, e.g. "Assumptions!$H$7"
ROW: dict[str, tuple[str, int]] = {}  # key -> (sheet, row) for rows spanning months or years


def mcol(m: int) -> str:
    return get_column_letter(MC0 + m - 1)


def ycol(y: int) -> str:
    return get_column_letter(YC0 + y - 1)


MLAST = mcol(MONTHS)
YLAST = ycol(YEARS)


def R(key: str) -> str:
    return REF[key]


def mcell(key: str, m: int) -> str:
    sheet, row = ROW[key]
    return f"{sheet}!{mcol(m)}{row}"


def mrange(key: str) -> str:
    sheet, row = ROW[key]
    return f"{sheet}!$D${row}:${MLAST}${row}"


def yrange(key: str) -> str:
    sheet, row = ROW[key]
    return f"{sheet}!$D${row}:${YLAST}${row}"


def put(ws, ref: str, value, font=F_BASE, fmt: str | None = None, fill=None, bold=False):
    cell = ws[ref]
    cell.value = value
    cell.font = Font(name=FONT, size=10, bold=True, color=font.color) if bold else font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def header_row(ws, row: int, labels: list[str], start_col: int = 1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=text)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def section(ws, row: int, text: str, last_col: int):
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).fill = FILL_SECTION
    c = ws.cell(row=row, column=2, value=text)
    c.font = F_BOLD


def title(ws, text: str, subtitle: str):
    put(ws, "A1", text, F_TITLE)
    put(ws, "A2", subtitle, F_NOTE)


def widths(ws, spec: dict[str, float]):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def is_cross_sheet(formula: str) -> bool:
    return "!" in formula


def fx(ws, ref: str, formula: str, fmt: str = NUM, bold=False, fill=None):
    """Write a formula; links to other sheets are green, local formulas black."""
    font = F_LINK if is_cross_sheet(formula) and formula.count("(") == 0 and formula.count("*") == 0 else F_BASE
    return put(ws, ref, formula, font, fmt, fill, bold)


def month_rows(ws, key: str, row: int, label: str, unit: str, formula_for, fmt: str = NUM, bold=False, fill=None):
    """A row of MONTHS formulas. formula_for(m) returns the formula for month m."""
    ROW[key] = (ws.title, row)
    put(ws, f"A{row}", key, F_NOTE)
    put(ws, f"B{row}", label, F_BOLD if bold else F_BASE)
    put(ws, f"C{row}", unit, F_NOTE)
    for m in range(1, MONTHS + 1):
        fx(ws, f"{mcol(m)}{row}", formula_for(m), fmt, bold, fill)


def year_rows(ws, key: str, row: int, label: str, unit: str, formula_for, fmt: str = NUM, bold=False, fill=None):
    ROW[key] = (ws.title, row)
    put(ws, f"A{row}", key, F_NOTE)
    put(ws, f"B{row}", label, F_BOLD if bold else F_BASE)
    put(ws, f"C{row}", unit, F_NOTE)
    for y in range(1, YEARS + 1):
        fx(ws, f"{ycol(y)}{row}", formula_for(y), fmt, bold, fill)


def timeline(ws):
    """Rows 3–6 on a monthly sheet: month number, farm year, calendar month, calendar year/date."""
    put(ws, "B3", "Month", F_BOLD)
    put(ws, "B4", "Farm year", F_BOLD)
    put(ws, "B5", "Calendar month", F_BOLD)
    put(ws, "B6", "Month starting", F_BOLD)
    sm, sy = R("start_month"), R("start_year")
    for m in range(1, MONTHS + 1):
        c = mcol(m)
        put(ws, f"{c}3", m, F_BASE, MONTHFMT, FILL_SECTION, bold=True)
        put(ws, f"{c}4", f"=INT(({c}$3-1)/12)+1", F_BASE, YEARFMT)
        put(ws, f"{c}5", f"=MOD({sm}+{c}$3-2,12)+1", F_BASE, "0")
        put(ws, f"{c}6", f"=DATE({sy}+INT(({sm}+{c}$3-2)/12),{c}$5,1)", F_BASE, DATEFMT)
    ws.freeze_panes = "D7"
    widths(ws, {"A": 12, "B": 44, "C": 14})
    for m in range(1, MONTHS + 1):
        ws.column_dimensions[mcol(m)].width = 11


def year_header(ws, row: int = 3):
    put(ws, f"B{row}", "Farm year", F_BOLD)
    for y in range(1, YEARS + 1):
        put(ws, f"{ycol(y)}{row}", y, F_BASE, YEARFMT, FILL_SECTION, bold=True)
    widths(ws, {"A": 12, "B": 46, "C": 16})
    for y in range(1, YEARS + 1):
        ws.column_dimensions[ycol(y)].width = 14


# ---------------------------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------------------------


def build_cover(ws):
    title(ws, "Farm financial model — structure (Stage 1)", f"Generated {date.today().isoformat()} by scripts/build_financial_model.py")
    lines = [
        ("Status", "STRUCTURE AND PLACEHOLDERS. Inputs labelled A or U are placeholders, not findings. Do not quote outputs as forecasts."),
        ("Purpose", "Tests whether each farm model fits the founder's capital, when cash turns positive, and which assumptions matter most."),
        ("Terms", "Real terms (constant 2026 RWF, no inflation). Month 1 = first month of spending. A farm year = 12 model months."),
        ("How to use", "Pick a scenario and a farm model on Control. Edit only blue cells. Results are on Metrics; detail on Annual and CashFlow."),
        ("Legend: blue", "Input you may edit (Control, Assumptions, Models, Curves, Sensitivity)."),
        ("Legend: black", "Formula. Do not overwrite."),
        ("Legend: green", "Direct link to another sheet."),
        ("Legend: yellow fill", "Key inputs to confirm with the founder."),
        ("Labels", "V verified (source ID) · E estimate (method shown) · A assumption · U unknown. See CLAUDE.md."),
        ("Structural constants", f"Sale age {SALE_AGE_M} months · first-birth lag {FIRST_BIRTH_LAG_M} months · receivable lag {RECEIVABLE_LAG_M} month. Change them in the script and rebuild."),
        ("Not modelled yet", "Loans and interest, inflation, tax losses carried forward, terminal value of orchards, value-added processing, agritourism."),
        ("Sheets", "Control · Assumptions · Models · Curves · Sensitivity · Land · Trees · Livestock · Feed · CapEx · CashFlow · Annual · Metrics"),
    ]
    for i, (k, v) in enumerate(lines, start=4):
        put(ws, f"A{i}", k, F_BOLD)
        put(ws, f"B{i}", v)
        ws[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, {"A": 22, "B": 120})


def build_control(ws):
    title(ws, "Control", "The two switches and the founder's key inputs. Blue cells only.")
    header_row(ws, 4, ["Key", "Input", "Value", "Unit", "Note"])
    for i, (key, label, value, unit, note) in enumerate(CONTROL, start=5):
        put(ws, f"A{i}", key, F_NOTE)
        put(ws, f"B{i}", label)
        fmt = {"capital": NUM, "area": NUM2}.get(key, "0")
        put(ws, f"C{i}", value, F_INPUT, fmt, FILL_KEY)
        put(ws, f"D{i}", unit, F_NOTE)
        put(ws, f"E{i}", note, F_NOTE)
        REF[key] = f"Control!$C${i}"
    r = 5 + len(CONTROL) + 1
    put(ws, f"B{r}", "Selected scenario", F_BOLD)
    put(ws, f"C{r}", f'=CHOOSE({R("scenario")},"Downside","Base","Upside")')
    REF["scenario_name"] = f"Control!$C${r}"
    put(ws, f"B{r + 1}", "Selected farm model", F_BOLD)
    put(ws, f"C{r + 1}", f'=CHOOSE({R("model")},"M1 Conservative mixed","M2 Diversified commercial","M3 Tech-enabled integrated")')
    REF["model_name"] = f"Control!$C${r + 1}"
    put(ws, f"B{r + 2}", "Capital available in USD (at the Assumptions rate)", F_BOLD)
    # FX is registered later; filled in after Assumptions exist.
    REF["_capital_usd_cell"] = f"C{r + 2}"
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(R("scenario").split("!")[1].replace("$", ""))
    dv.add(R("model").split("!")[1].replace("$", ""))
    dvm = DataValidation(type="whole", operator="between", formula1="1", formula2="12")
    ws.add_data_validation(dvm)
    dvm.add(R("start_month").split("!")[1].replace("$", ""))
    widths(ws, {"A": 12, "B": 48, "C": 26, "D": 10, "E": 52})


def build_assumptions(ws):
    title(ws, "Assumptions", "Uncertain inputs with downside / base / upside values. Selected = the value for the scenario chosen on Control.")
    header_row(ws, 4, ["ID", "Group", "Parameter", "Unit", "Downside", "Base", "Upside", "Selected", "Label", "Source", "Notes"])
    rows = []
    for (aid, group, name, unit, d, b, u, label, source, note) in A:
        rows.append((aid, group, name, unit, d, b, u, label, source, note))
    for suffix, name, unit, gt, sh, label, note in LIVESTOCK:
        for code, species in SPECIES:
            vals = gt if code == "GT" else sh
            rows.append((f"{code}_{suffix}", f"Livestock: {species}", name, unit, *vals, label, PENDING, note))
    for key, cname, dens, surv, seed, plant, ymult, mkt, price, cycle, curve in CROPS:
        grp = f"Crop: {cname}"
        rows.append((f"{key}_DENS", grp, "Planting density", "plants per ha", dens, dens, dens, "A", PENDING, "Design choice; confirm spacing"))
        rows.append((f"{key}_SURV", grp, "Final stand after gap filling", "share", *surv, "A", PENDING, ""))
        rows.append((f"{key}_SEED", grp, "Seedling or grafted plant", "RWF each", *seed, "A", PENDING, ""))
        rows.append((f"{key}_PLANT", grp, "Planting (hole, compost, labour, trellis)", "RWF per plant", *plant, "A", PENDING, ""))
        rows.append((f"{key}_YMULT", grp, "Yield multiplier on the curve", "factor", *ymult, "A", PENDING, ""))
        rows.append((f"{key}_MKT", grp, "Marketable share of harvest", "share", *mkt, "A", PENDING, ""))
        rows.append((f"{key}_PRICE", grp, "Farm-gate price", "RWF per kg", *price, "A", PENDING, ""))
        rows.append((f"{key}_CYCLE", grp, "Replant cycle (0 = perennial)", "years", cycle, cycle, cycle, "A", PENDING, ""))
    for tid, name, unit, vals, label, note in TIMBER:
        rows.append((tid, "Boundary timber", name, unit, *vals, label, PENDING, note))

    r = 5
    last_group = None
    sc = R("scenario")
    for (aid, group, name, unit, d, b, u, label, source, note) in rows:
        if group != last_group:
            section(ws, r, group, 11)
            r += 1
            last_group = group
        fmt = PCT if unit.startswith("share") else (NUM2 if isinstance(b, float) and b < 100 else NUM)
        put(ws, f"A{r}", aid, F_NOTE)
        put(ws, f"B{r}", group, F_NOTE)
        put(ws, f"C{r}", name)
        put(ws, f"D{r}", unit, F_NOTE)
        for col, v in zip("EFG", (d, b, u)):
            put(ws, f"{col}{r}", v, F_INPUT, fmt)
        put(ws, f"H{r}", f"=CHOOSE({sc},E{r},F{r},G{r})", F_BASE, fmt, FILL_TOTAL)
        put(ws, f"I{r}", label, F_BOLD)
        put(ws, f"J{r}", source, F_NOTE)
        put(ws, f"K{r}", note, F_NOTE)
        REF[aid] = f"Assumptions!$H${r}"
        r += 1
    ws.freeze_panes = "E5"
    widths(ws, {"A": 14, "B": 22, "C": 46, "D": 22, "E": 12, "F": 12, "G": 12, "H": 12, "I": 7, "J": 36, "K": 44})


def build_models(ws):
    title(ws, "Models", "What defines each farm model. Selected = the column for the model chosen on Control.")
    header_row(ws, 4, ["ID", "Group", "Parameter", "Unit", "M1", "M2", "M3", "Selected", "Label", "Notes"])
    r = 5
    last_group = None
    mc = R("model")
    for (mid, group, name, unit, m1, m2, m3, label, note) in M:
        if group != last_group:
            section(ws, r, group, 10)
            r += 1
            last_group = group
        fmt = PCT if unit == "share" else (NUM1 if isinstance(m1, float) and m1 < 100 else NUM)
        put(ws, f"A{r}", mid, F_NOTE)
        put(ws, f"B{r}", group, F_NOTE)
        put(ws, f"C{r}", name)
        put(ws, f"D{r}", unit, F_NOTE)
        for col, v in zip("EFG", (m1, m2, m3)):
            put(ws, f"{col}{r}", v, F_INPUT, fmt)
        put(ws, f"H{r}", f"=CHOOSE({mc},E{r},F{r},G{r})", F_BASE, fmt, FILL_TOTAL)
        put(ws, f"I{r}", label, F_BOLD)
        put(ws, f"J{r}", note, F_NOTE)
        REF[mid] = f"Models!$H${r}"
        r += 1
    ws.freeze_panes = "E5"
    widths(ws, {"A": 12, "B": 30, "C": 52, "D": 16, "E": 12, "F": 12, "G": 12, "H": 12, "I": 7, "J": 44})


def build_curves(ws):
    title(ws, "Curves", "Yield per plant by farm year, harvest seasonality by calendar month, pasture growth by month.")
    header_row(ws, 4, ["ID", "Crop", "Unit"] + [f"Y{y}" for y in range(1, YEARS + 1)] + ["Label", "Source / note"])
    r = 5
    for key, cname, *_rest in CROPS:
        curve = _rest[-1]
        put(ws, f"A{r}", f"{key}_CURVE", F_NOTE)
        put(ws, f"B{r}", cname)
        put(ws, f"C{r}", "kg per plant", F_NOTE)
        for y, v in enumerate(curve, start=1):
            put(ws, f"{ycol(y)}{r}", v, F_INPUT, NUM1)
        put(ws, f"N{r}", "A", F_BOLD)
        put(ws, f"O{r}", PENDING + ". Planted in year 1; short-cycle crops repeat their first years.", F_NOTE)
        ROW[f"{key}_CURVE"] = ("Curves", r)
        r += 1
    r += 2
    sr = r
    header_row(ws, sr, ["ID", "Crop", "Unit"] + MONTH_NAMES + ["Sum (must be 1)", "Label", "Note"])
    r = sr + 1
    for key, cname, *_ in CROPS:
        put(ws, f"A{r}", f"{key}_SEASON", F_NOTE)
        put(ws, f"B{r}", cname)
        put(ws, f"C{r}", "share of year", F_NOTE)
        for i, v in enumerate(SEASON[key]):
            put(ws, f"{get_column_letter(4 + i)}{r}", v, F_INPUT, PCT)
        put(ws, f"P{r}", f"=SUM(D{r}:O{r})", F_BASE, PCT)
        put(ws, f"Q{r}", "A", F_BOLD)
        put(ws, f"R{r}", "Uniform until the Rwandan harvest calendar is sourced", F_NOTE)
        ROW[f"{key}_SEASON"] = ("Curves", r)
        r += 1
    r += 2
    header_row(ws, r, ["ID", "Series", "Unit"] + MONTH_NAMES + ["Average (must be 1)", "Label", "Note"])
    r += 1
    put(ws, f"A{r}", "PGROWTH", F_NOTE)
    put(ws, f"B{r}", "Pasture growth factor")
    put(ws, f"C{r}", "x average", F_NOTE)
    for i, v in enumerate(PASTURE_GROWTH):
        put(ws, f"{get_column_letter(4 + i)}{r}", v, F_INPUT, NUM2)
    put(ws, f"P{r}", f"=AVERAGE(D{r}:O{r})", F_BASE, NUM2)
    put(ws, f"Q{r}", "A", F_BOLD)
    put(ws, f"R{r}", "Shape follows the bimodal rains; calibrate to zone rainfall", F_NOTE)
    ROW["PGROWTH"] = ("Curves", r)
    widths(ws, {"A": 14, "B": 26, "C": 14, "N": 7, "O": 60, "P": 16, "Q": 7, "R": 50})
    for col in range(4, 16):
        ws.column_dimensions[get_column_letter(col)].width = 9


def build_sensitivity(ws):
    title(ws, "Sensitivity", "Multipliers that flex the drivers that matter most. Leave at 1 (delay 0) for the plain scenario.")
    header_row(ws, 4, ["Key", "Driver", "Setting", "How it is applied"])
    r = 5
    for key, driver, default, how in SENS:
        put(ws, f"A{r}", key, F_NOTE)
        put(ws, f"B{r}", driver)
        put(ws, f"C{r}", default, F_INPUT, NUM2 if key != "S_DELAY" else "0")
        put(ws, f"D{r}", how, F_NOTE)
        REF[key] = f"Sensitivity!$C${r}"
        r += 1
    r += 2
    put(ws, f"B{r}", "Results log (Stage 3): set one driver at a time to 0.8 and 1.2, record the outputs, reset to 1.", F_BOLD)
    r += 1
    header_row(ws, r, ["", "Driver", "Setting", "NPV (RWF)", "Funding gap (RWF)", "Payback year", "Date / note"], start_col=1)
    widths(ws, {"A": 16, "B": 28, "C": 12, "D": 60, "E": 18, "F": 14, "G": 30})


def build_land(ws):
    title(ws, "Land", "Areas by zone and crop for the selected model.")
    header_row(ws, 3, ["Key", "Item", "Value", "Unit"])
    rows = [
        ("L_AREA", "Total area", f"={R('area')}", "ha", NUM2),
        ("L_PERIM", "Perimeter of a square of that area", f"=4*SQRT(C4*10000)", "m", NUM),
        ("L_BOUND", "Boundary length with shape factor", f"=C5*{R('T_SHAPE')}", "m", NUM),
        ("L_Z01", "Z0+Z1 base, store, nursery", f"=C4*{R('Z01')}", "ha", NUM2),
        ("L_Z2", "Z2 orchard", f"=C4*{R('Z2')}", "ha", NUM2),
        ("L_Z3", "Z3 paddocks and fodder bank", f"=C4*{R('Z3')}", "ha", NUM2),
        ("L_Z4", "Z4 boundary and shelter trees", f"=C4*{R('Z4')}", "ha", NUM2),
        ("L_Z5", "Z5 buffer and conservation", f"=C4*{R('Z5')}", "ha", NUM2),
        ("L_PADDOCK", "Paddocks (grazed)", f"=C9*(1-{R('FODDER')})", "ha", NUM2),
        ("L_FODDER", "Fodder bank (cut-and-carry)", f"=C9*{R('FODDER')}", "ha", NUM2),
    ]
    r = 4
    for key, label, formula, unit, fmt in rows:
        put(ws, f"A{r}", key, F_NOTE)
        put(ws, f"B{r}", label)
        fx(ws, f"C{r}", formula, fmt)
        put(ws, f"D{r}", unit, F_NOTE)
        REF[key] = f"Land!$C${r}"
        r += 1
    r += 1
    section(ws, r, "Orchard by crop", 4)
    r += 1
    for key, cname, *_ in CROPS:
        put(ws, f"A{r}", f"L_{key}", F_NOTE)
        put(ws, f"B{r}", cname)
        fx(ws, f"C{r}", f"={R('L_Z2')}*{R('OR_' + key)}", NUM2)
        put(ws, f"D{r}", "ha", F_NOTE)
        REF[f"L_{key}"] = f"Land!$C${r}"
        r += 1
    r += 1
    section(ws, r, "Checks", 4)
    r += 1
    put(ws, f"B{r}", "Land shares sum to 100%")
    fx(ws, f"C{r}", f"={R('Z01')}+{R('Z2')}+{R('Z3')}+{R('Z4')}+{R('Z5')}", PCT)
    REF["CHK_LAND"] = f"Land!$C${r}"
    r += 1
    put(ws, f"B{r}", "Orchard mix sums to 100%")
    fx(ws, f"C{r}", "=" + "+".join(R("OR_" + c[0]) for c in CROPS), PCT)
    REF["CHK_ORCH"] = f"Land!$C${r}"
    widths(ws, {"A": 12, "B": 44, "C": 14, "D": 8})


def build_trees(ws):
    title(ws, "Trees", "Annual orchard and boundary-tree engine for the selected model and scenario.")
    year_header(ws, 3)
    r = 5
    yrow = 3
    est_cells, gap_cells, rev_rows, prod_rows, replant_rows = [], [], [], [], []
    for key, cname, *_ in CROPS:
        section(ws, r, cname, 3 + YEARS)
        r += 1
        a = lambda s: R(f"{key}_{s}")  # noqa: E731
        singles = [
            (f"{key}_AREA", "Area", f"={R('L_' + key)}", "ha", NUM2),
            (f"{key}_PLANTS", "Plants planted", f"=ROUND(C{r}*{a('DENS')},0)", "plants", NUM),
            (f"{key}_STAND", "Final stand", f"=C{r + 1}*{a('SURV')}", "plants", NUM),
            (f"{key}_EST", "Establishment cost", f"=C{r + 1}*({a('SEED')}+{a('PLANT')})*{R('S_CAPEX')}", "RWF", NUM),
            (f"{key}_GAP", "Gap-filling cost (year 2)", f"=C{r + 1}*(1-{a('SURV')})*{a('SEED')}*{R('S_CAPEX')}", "RWF", NUM),
        ]
        for i, (k, label, formula, unit, fmt) in enumerate(singles):
            put(ws, f"A{r + i}", k, F_NOTE)
            put(ws, f"B{r + i}", label)
            fx(ws, f"C{r + i}", formula, fmt)
            REF[k] = f"Trees!$C${r + i}"
        est_cells.append(REF[f"{key}_EST"])
        gap_cells.append(REF[f"{key}_GAP"])
        r += len(singles)
        curve = ROW[f"{key}_CURVE"][1]
        cyc, delay = a("CYCLE"), R("S_DELAY")
        year_rows(ws, f"{key}_AGE", r, "Age in production cycle", "years",
                  lambda y: f"=IF({cyc}=0,{ycol(y)}${yrow}-{delay},MOD({ycol(y)}${yrow}-1,MAX(1,{cyc}))+1)", "0")
        year_rows(ws, f"{key}_YLD", r + 1, "Yield per plant", "kg",
                  lambda y: f"=IF({ycol(y)}{r}>=1,INDEX(Curves!$D${curve}:${YLAST}${curve},1,MIN({YEARS},MAX(1,{ycol(y)}{r}))),0)*{a('YMULT')}*{R('S_YIELD')}", NUM1)
        year_rows(ws, f"{key}_PROD", r + 2, "Marketable production", "kg",
                  lambda y: f"={R(key + '_STAND')}*{ycol(y)}{r + 1}*{a('MKT')}", NUM)
        year_rows(ws, f"{key}_REV", r + 3, "Sales", "RWF",
                  lambda y: f"={ycol(y)}{r + 2}*{a('PRICE')}*{R('S_PRICE_FRUIT')}", NUM)
        year_rows(ws, f"{key}_REPL", r + 4, "Replanting (short-cycle crops)", "RWF",
                  lambda y: f"=IF(AND({cyc}>0,{ycol(y)}${yrow}>1,MOD({ycol(y)}${yrow}-1,MAX(1,{cyc}))=0),{R(key + '_PLANTS')}*({a('SEED')}+{a('PLANT')})*{R('S_CAPEX')},0)", NUM)
        rev_rows.append(ROW[f"{key}_REV"][1])
        prod_rows.append(ROW[f"{key}_PROD"][1])
        replant_rows.append(ROW[f"{key}_REPL"][1])
        r += 6

    section(ws, r, "Boundary timber and poles", 3 + YEARS)
    r += 1
    t = lambda s: R(s)  # noqa: E731
    singles = [
        ("TIM_TREES", "Trees planted on the boundary", f"=ROUND({R('L_BOUND')}/{t('T_SPACING')},0)", "trees", NUM),
        ("TIM_STAND", "Final stand", f"=C{r}*{t('T_SURV')}", "trees", NUM),
        ("TIM_EST", "Establishment cost", f"=C{r}*({t('T_SEED')}+{t('T_PLANT')})*{R('S_CAPEX')}", "RWF", NUM),
        ("TIM_MEMO", "Memo: standing value of remaining trees at Y10 (not in NPV)", f"=C{r + 1}*(1-{t('T_POLESHARE')})*{t('T_STAND')}", "RWF", NUM),
    ]
    for i, (k, label, formula, unit, fmt) in enumerate(singles):
        put(ws, f"A{r + i}", k, F_NOTE)
        put(ws, f"B{r + i}", label)
        fx(ws, f"C{r + i}", formula, fmt)
        REF[k] = f"Trees!$C${r + i}"
    r += len(singles)
    year_rows(ws, "TIM_REV", r, "Pole sales (first thinning)", "RWF",
              lambda y: f"=IF({ycol(y)}${yrow}={t('T_POLEYEAR')},{R('TIM_STAND')}*{t('T_POLESHARE')}*{t('T_POLEPRICE')},0)", NUM)
    r += 2

    section(ws, r, "Totals", 3 + YEARS)
    r += 1
    put(ws, f"A{r}", "TR_EST", F_NOTE)
    put(ws, f"B{r}", "Orchard establishment cost (all crops)", F_BOLD)
    fx(ws, f"C{r}", "=" + "+".join(est_cells), NUM, bold=True)
    REF["TR_EST"] = f"Trees!$C${r}"
    r += 1
    put(ws, f"A{r}", "TR_GAP", F_NOTE)
    put(ws, f"B{r}", "Gap-filling cost, year 2 (all crops)", F_BOLD)
    fx(ws, f"C{r}", "=" + "+".join(gap_cells), NUM, bold=True)
    REF["TR_GAP"] = f"Trees!$C${r}"
    r += 1
    year_rows(ws, "TR_REV", r, "Fruit sales (all crops)", "RWF",
              lambda y: "=" + "+".join(f"{ycol(y)}{x}" for x in rev_rows), NUM, bold=True)
    year_rows(ws, "TR_PROD", r + 1, "Marketable fruit (all crops)", "kg",
              lambda y: "=" + "+".join(f"{ycol(y)}{x}" for x in prod_rows), NUM, bold=True)
    year_rows(ws, "TR_REPL", r + 2, "Replanting (all crops)", "RWF",
              lambda y: "=" + "+".join(f"{ycol(y)}{x}" for x in replant_rows), NUM, bold=True)
    ws.freeze_panes = "D4"


def build_livestock(ws):
    title(ws, "Livestock", f"Monthly herd engine. Young are sold at {SALE_AGE_M} months; females give first birth {FIRST_BIRTH_LAG_M} months after joining.")
    timeline(ws)
    r = 8
    for code, species in SPECIES:
        section(ws, r, species, 3 + MONTHS)
        r += 1
        p = lambda s: R(f"{code}_{s}")  # noqa: E731
        m_ = lambda s: R(f"{code}_{s}")  # noqa: E731
        rows = {}
        names = ["FBUY", "MBUYT", "FOPEN", "FDEATH", "FCULL", "BIRTHS", "YSURV", "YREACH", "RETAINED",
                 "FCLOSE", "YSOLD", "YHAND", "MNEED", "MALES", "MBUY", "HEAD", "SOLDHD", "REV_Y", "REV_C",
                 "BUYCOST", "VALUE", "DEMAND"]
        for i, n in enumerate(names):
            rows[n] = r + i
        c = mcol
        t1m, t2m = m_("T1M"), m_("T2M")

        def prev(n, m):
            return f"{c(m - 1)}{rows[n]}" if m > 1 else "0"

        month_rows(ws, f"{code}_FBUY", rows["FBUY"], "Breeding females bought", "head",
                   lambda m: f"=IF({c(m)}$3={t1m},{m_('T1F')},0)+IF({c(m)}$3={t2m},{m_('T2F')},0)", NUM1)
        month_rows(ws, f"{code}_MBUYT", rows["MBUYT"], "Breeding males bought in tranches", "head",
                   lambda m: f"=IF({c(m)}$3={t1m},{m_('T1MALE')},0)+IF({c(m)}$3={t2m},{m_('T2MALE')},0)", NUM1)
        month_rows(ws, f"{code}_FOPEN", rows["FOPEN"], "Breeding females, opening", "head",
                   lambda m: f"={prev('FCLOSE', m)}", NUM1)
        month_rows(ws, f"{code}_FDEATH", rows["FDEATH"], "Female deaths", "head",
                   lambda m: f"={c(m)}{rows['FOPEN']}*MIN(1,{p('AMORT')}*{R('S_MORT')})/12", NUM2)
        month_rows(ws, f"{code}_FCULL", rows["FCULL"], "Females culled (sold)", "head",
                   lambda m: f"={c(m)}{rows['FOPEN']}*{p('CULL')}/12", NUM2)
        month_rows(ws, f"{code}_BIRTHS", rows["BIRTHS"], "Young born", "head",
                   lambda m: (f"={c(m - FIRST_BIRTH_LAG_M)}{rows['FCLOSE']}*{p('PART')}/12*{p('LIT')}"
                              if m > FIRST_BIRTH_LAG_M else "=0"), NUM2)
        month_rows(ws, f"{code}_YSURV", rows["YSURV"], "Young surviving to weaning", "head",
                   lambda m: f"={c(m)}{rows['BIRTHS']}*(1-MIN(1,{p('YMORT')}*{R('S_MORT')}))", NUM2)
        month_rows(ws, f"{code}_YREACH", rows["YREACH"], f"Young reaching sale age ({SALE_AGE_M} months)", "head",
                   lambda m: f"={c(m - SALE_AGE_M)}{rows['YSURV']}" if m > SALE_AGE_M else "=0", NUM2)
        month_rows(ws, f"{code}_RETAINED", rows["RETAINED"], "Female young kept as replacements", "head",
                   lambda m: (f"=MIN({c(m)}{rows['YREACH']}*{p('FSHARE')}*{R('RETAIN')},"
                              f"MAX(0,{m_('TARGET')}-({c(m)}{rows['FOPEN']}-{c(m)}{rows['FDEATH']}-{c(m)}{rows['FCULL']}+{c(m)}{rows['FBUY']})))"), NUM2)
        month_rows(ws, f"{code}_FCLOSE", rows["FCLOSE"], "Breeding females, closing", "head",
                   lambda m: (f"={c(m)}{rows['FOPEN']}-{c(m)}{rows['FDEATH']}-{c(m)}{rows['FCULL']}"
                              f"+{c(m)}{rows['FBUY']}+{c(m)}{rows['RETAINED']}"), NUM1, bold=True)
        month_rows(ws, f"{code}_YSOLD", rows["YSOLD"], "Young sold", "head",
                   lambda m: f"={c(m)}{rows['YREACH']}-{c(m)}{rows['RETAINED']}", NUM2)
        month_rows(ws, f"{code}_YHAND", rows["YHAND"], "Young on hand", "head",
                   lambda m: f"={prev('YHAND', m)}+{c(m)}{rows['YSURV']}-{c(m)}{rows['YREACH']}", NUM1)
        month_rows(ws, f"{code}_MNEED", rows["MNEED"], "Breeding males needed", "head",
                   lambda m: f"=IF({c(m)}{rows['FCLOSE']}>0,ROUNDUP({c(m)}{rows['FCLOSE']}/{p('RATIO')},0),0)", "0")
        month_rows(ws, f"{code}_MALES", rows["MALES"], "Breeding males", "head",
                   lambda m: f"=MAX({prev('MALES', m)}+{c(m)}{rows['MBUYT']},{c(m)}{rows['MNEED']})", "0")
        month_rows(ws, f"{code}_MBUY", rows["MBUY"], "Breeding males bought (all)", "head",
                   lambda m: f"={c(m)}{rows['MALES']}-{prev('MALES', m)}", "0")
        month_rows(ws, f"{code}_HEAD", rows["HEAD"], "All animals", "head",
                   lambda m: f"={c(m)}{rows['FCLOSE']}+{c(m)}{rows['MALES']}+{c(m)}{rows['YHAND']}", NUM1, bold=True)
        month_rows(ws, f"{code}_SOLDHD", rows["SOLDHD"], "Animals sold", "head",
                   lambda m: f"={c(m)}{rows['YSOLD']}+{c(m)}{rows['FCULL']}", NUM2)
        month_rows(ws, f"{code}_REV_Y", rows["REV_Y"], "Sales: young stock", "RWF",
                   lambda m: (f"={c(m)}{rows['YSOLD']}*{p('SALEWT')}*{p('PRICE')}*{R('S_PRICE_LIVE')}"
                              f"*(1+{R('BSHARE')}*{p('BPREM')})"), NUM)
        month_rows(ws, f"{code}_REV_C", rows["REV_C"], "Sales: culled adults", "RWF",
                   lambda m: f"={c(m)}{rows['FCULL']}*{p('ADULTWT')}*{p('PRICE')}*{R('S_PRICE_LIVE')}", NUM)
        month_rows(ws, f"{code}_BUYCOST", rows["BUYCOST"], "Animal purchases", "RWF",
                   lambda m: (f"=({c(m)}{rows['FBUY']}*{p('FPRICE')}+{c(m)}{rows['MBUY']}*{p('MPRICE')})*{R('S_CAPEX')}"), NUM)
        month_rows(ws, f"{code}_VALUE", rows["VALUE"], "Breeding herd value (insurance basis)", "RWF",
                   lambda m: f"={c(m)}{rows['FCLOSE']}*{p('FPRICE')}+{c(m)}{rows['MALES']}*{p('MPRICE')}", NUM)
        month_rows(ws, f"{code}_DEMAND", rows["DEMAND"], "Feed demand", "kg DM",
                   lambda m: (f"=(({c(m)}{rows['FCLOSE']}+{c(m)}{rows['MALES']})*{p('DMI_A')}"
                              f"+{c(m)}{rows['YHAND']}*{p('DMI_Y')})*{DAYS_PER_MONTH}"), NUM)
        r += len(names) + 1


def build_feed(ws):
    title(ws, "Feed", "Monthly dry-matter balance: forage grown, conserved and bought.")
    timeline(ws)
    c = mcol
    pm = R("past_m")
    growth = ROW["PGROWTH"][1]
    month_rows(ws, "FD_DEMAND", 8, "Feed demand (all animals)", "kg DM",
               lambda m: f"={mcell('GT_DEMAND', m)}+{mcell('SH_DEMAND', m)}", NUM)
    month_rows(ws, "FD_READY", 9, "Forage established (1 = yes)", "flag",
               lambda m: f"=IF({c(m)}$3>={pm}+{R('READY')},1,0)", "0")
    month_rows(ws, "FD_SUPPLY", 10, "Forage available to animals", "kg DM",
               lambda m: (f"=({R('L_PADDOCK')}*{R('PAST_Y')}*1000*{R('GRAZE')}+{R('L_FODDER')}*{R('FOD_Y')}*1000*{R('CUT')})/12"
                          f"*INDEX(Curves!$D${growth}:$O${growth},1,{c(m)}$5)*{c(m)}9"), NUM)
    month_rows(ws, "FD_SURPLUS", 11, "Surplus forage", "kg DM", lambda m: f"=MAX(0,{c(m)}10-{c(m)}8)", NUM)
    month_rows(ws, "FD_DEFICIT", 12, "Forage deficit", "kg DM", lambda m: f"=MAX(0,{c(m)}8-{c(m)}10)", NUM)
    month_rows(ws, "FD_CONS", 13, "Conserved as hay/silage", "kg DM",
               lambda m: f"={c(m)}11*{R('CONSERVE')}*{R('CONS_EFF')}", NUM)
    month_rows(ws, "FD_AVAIL", 14, "Conserved feed available", "kg DM",
               lambda m: f"={(c(m - 1) + '16') if m > 1 else '0'}+{c(m)}13", NUM)
    month_rows(ws, "FD_DRAWN", 15, "Conserved feed used", "kg DM", lambda m: f"=MIN({c(m)}14,{c(m)}12)", NUM)
    month_rows(ws, "FD_STORE", 16, "Conserved feed in store, closing", "kg DM", lambda m: f"={c(m)}14-{c(m)}15", NUM)
    month_rows(ws, "FD_BUY", 17, "Feed bought", "kg DM", lambda m: f"={c(m)}12-{c(m)}15", NUM, bold=True)
    month_rows(ws, "FD_BUYCOST", 18, "Feed bought", "RWF", lambda m: f"={c(m)}17*{R('FEED_COST')}*{R('S_FEED')}", NUM, bold=True)
    month_rows(ws, "FD_CONSCOST", 19, "Conservation cost", "RWF", lambda m: f"={c(m)}13*{R('CONS_COST')}", NUM)


CAPEX = [
    # (id, item, category, essential?, qty formula, unit, unit-cost formula, month formula, life yrs, cost already includes S_CAPEX?)
    ("CX01", "Soil tests", "Studies", "Essential", lambda: f"={R('SOIL_N')}", "samples", lambda: f"={R('SOIL')}", lambda: "=1", 5, False),
    ("CX02", "Water tests", "Studies", "Essential", lambda: f"={R('WATER_N')}", "samples", lambda: f"={R('WATER')}", lambda: "=1", 5, False),
    ("CX03", "Survey, title and legal checks", "Studies", "Essential", lambda: "=1", "lump", lambda: f"={R('SURVEY')}", lambda: "=1", 10, False),
    ("CX04", "Permits and compliance", "Studies", "Essential", lambda: "=1", "lump", lambda: f"={R('PERMIT')}", lambda: "=2", 10, False),
    ("CX05", "Professional set-up visits", "Studies", "Essential", lambda: "=1", "lump", lambda: f"={R('PROF')}", lambda: "=2", 5, False),
    ("CX06", "Perimeter fence", "Fencing", "Essential", lambda: f"={R('L_BOUND')}", "m", lambda: f"={R('FENCE')}", lambda: "=3", 10, False),
    ("CX07", "Internal paddock fence", "Fencing", "Essential", lambda: f"={R('L_Z3')}*{R('IFENCE_M')}", "m", lambda: f"={R('IFENCE')}", lambda: f"=MAX(1,{R('GT_T1M')}-2)", 8, False),
    ("CX08", "Gates", "Fencing", "Essential", lambda: f"={R('GATES')}", "count", lambda: f"={R('GATE')}", lambda: "=3", 10, False),
    ("CX09", "Water storage", "Water", "Essential", lambda: f"={R('STORAGE')}", "m3", lambda: f"={R('TANK')}", lambda: "=3", 15, False),
    ("CX10", "Rainwater harvesting fittings", "Water", "Essential", lambda: "=1", "lump", lambda: f"={R('GUTTER')}", lambda: "=3", 10, False),
    ("CX11", "Borehole", "Water", "Optional", lambda: f"={R('BOREHOLE')}", "flag", lambda: f"={R('BORE')}", lambda: "=4", 20, False),
    ("CX12", "Solar pump system", "Water", "Optional", lambda: f"={R('SOLARPUMP')}", "flag", lambda: f"={R('PUMP')}", lambda: "=4", 10, False),
    ("CX13", "Drip irrigation on the orchard", "Water", "Optional", lambda: f"={R('L_Z2')}*{R('DRIPON')}", "ha", lambda: f"={R('DRIP')}", lambda: f"=MAX(1,{R('plant_m')}-1)", 8, False),
    ("CX14", "Orchard establishment (plants and planting)", "Orchard", "Essential", lambda: "=1", "lump", lambda: f"={R('TR_EST')}", lambda: f"={R('plant_m')}", 20, True),
    ("CX15", "Orchard gap filling", "Orchard", "Essential", lambda: "=1", "lump", lambda: f"={R('TR_GAP')}", lambda: f"={R('plant_m')}+12", 20, True),
    ("CX16", "Boundary trees", "Orchard", "Essential", lambda: "=1", "lump", lambda: f"={R('TIM_EST')}", lambda: f"={R('plant_m')}", 20, True),
    ("CX17", "Pasture establishment", "Feed", "Essential", lambda: f"={R('L_PADDOCK')}", "ha", lambda: f"={R('PAST_EST')}", lambda: f"={R('past_m')}", 5, False),
    ("CX18", "Fodder bank establishment", "Feed", "Essential", lambda: f"={R('L_FODDER')}", "ha", lambda: f"={R('FOD_EST')}", lambda: f"={R('past_m')}", 5, False),
    ("CX19", "Animal house", "Livestock", "Essential", lambda: f"=({R('GT_TARGET')}+{R('SH_TARGET')})*1.1*{R('M2HEAD')}", "m2", lambda: f"={R('HOUSE')}", lambda: f"=MAX(1,{R('GT_T1M')}-1)", 15, False),
    ("CX20", "Quarantine pen and handling race", "Livestock", "Essential", lambda: "=1", "lump", lambda: f"={R('QUAR')}", lambda: f"=MAX(1,{R('GT_T1M')}-1)", 10, False),
    ("CX21", "Feed store", "Livestock", "Essential", lambda: "=1", "lump", lambda: f"={R('FEEDSTORE')}", lambda: f"=MAX(1,{R('GT_T1M')}-1)", 15, False),
    ("CX22", "Livestock weighing scale", "Livestock", "Essential", lambda: "=1", "lump", lambda: f"={R('SCALE')}", lambda: f"={R('GT_T1M')}", 8, False),
    ("CX23", "Tools and small equipment", "Equipment", "Essential", lambda: "=1", "lump", lambda: f"={R('TOOLS')}", lambda: "=2", 5, False),
    ("CX24", "Solar PV with battery", "Energy and tech", "Essential", lambda: f"={R('PV_KWP')}", "kWp", lambda: f"={R('PV')}", lambda: "=3", 10, False),
    ("CX25", "Connectivity", "Energy and tech", "Essential", lambda: "=1", "lump", lambda: f"={R('CONNECT')}", lambda: "=3", 5, False),
    ("CX26", "Sensors and controllers", "Energy and tech", "Optional", lambda: "=1", "lump", lambda: f"={R('SENSORS')}", lambda: "=6", 5, False),
    ("CX27", "Software and devices", "Energy and tech", "Optional", lambda: "=1", "lump", lambda: f"={R('SOFTWARE')}", lambda: "=6", 3, False),
    ("CX28", "Security (lights, cameras)", "Energy and tech", "Essential", lambda: "=1", "lump", lambda: f"={R('SECURITY')}", lambda: "=3", 7, False),
    ("CX29", "Nursery", "Orchard", "Optional", lambda: "=1", "lump", lambda: f"={R('NURSERY')}", lambda: "=2", 5, False),
    ("CX30", "Cool store / packhouse", "Post-harvest", "Optional", lambda: "=1", "lump", lambda: f"={R('COOLSTORE')}", lambda: "=24", 15, False),
]


def build_capex(ws):
    title(ws, "CapEx", "Capital items for the selected model. Contingency is added on top. Animal purchases are on Livestock.")
    header_row(ws, 4, ["ID", "Item", "Category", "Essential / optional", "Quantity", "Unit", "Unit cost (RWF)",
                       "Total (RWF)", "Month incurred", "Year incurred", "Life (yrs)", "Depreciation per yr (RWF)"])
    first = 5
    for i, (cid, item, cat, ess, qty, unit, cost, month, life, prepriced) in enumerate(CAPEX):
        r = first + i
        put(ws, f"A{r}", cid, F_NOTE)
        put(ws, f"B{r}", item)
        put(ws, f"C{r}", cat, F_NOTE)
        put(ws, f"D{r}", ess)
        fx(ws, f"E{r}", qty(), NUM2)
        put(ws, f"F{r}", unit, F_NOTE)
        fx(ws, f"G{r}", cost(), NUM)
        total = f"=E{r}*G{r}" if prepriced else f"=E{r}*G{r}*{R('S_CAPEX')}"
        fx(ws, f"H{r}", total, NUM)
        fx(ws, f"I{r}", month(), "0")
        fx(ws, f"J{r}", f"=INT((I{r}-1)/12)+1", YEARFMT)
        put(ws, f"K{r}", life, F_INPUT, "0")
        fx(ws, f"L{r}", f"=IF(K{r}>0,H{r}/K{r},0)", NUM)
    last = first + len(CAPEX) - 1
    REF["CX_FIRST"], REF["CX_LAST"] = str(first), str(last)
    r = last + 2
    totals = [
        ("CX_ESS", "Essential items", f'=SUMIF($D${first}:$D${last},"Essential",$H${first}:$H${last})'),
        ("CX_OPT", "Optional items", f'=SUMIF($D${first}:$D${last},"Optional",$H${first}:$H${last})'),
        ("CX_SUB", "All items", f"=SUM($H${first}:$H${last})"),
        ("CX_CONT", "Contingency", f"=H{r + 2}*{R('CONT')}"),
        ("CX_TOT", "Capital items including contingency", f"=H{r + 2}+H{r + 3}"),
    ]
    for i, (k, label, formula) in enumerate(totals):
        put(ws, f"A{r + i}", k, F_NOTE)
        put(ws, f"B{r + i}", label, F_BOLD)
        fx(ws, f"H{r + i}", formula, NUM, bold=True, fill=FILL_TOTAL)
        REF[k] = f"CapEx!$H${r + i}"
    ws.freeze_panes = "C5"
    widths(ws, {"A": 8, "B": 42, "C": 16, "D": 12, "E": 11, "F": 9, "G": 15, "H": 15, "I": 10, "J": 10, "K": 8, "L": 16})


def depreciation(year_ref: str) -> str:
    """Straight-line depreciation for one farm year, from the CapEx sheet, including contingency."""
    f, l = REF["CX_FIRST"], REF["CX_LAST"]
    return (f"SUMPRODUCT((CapEx!$J${f}:$J${l}<={year_ref})*(CapEx!$J${f}:$J${l}+CapEx!$K${f}:$K${l}>{year_ref})"
            f"*CapEx!$L${f}:$L${l})*(1+{R('CONT')})")


def build_cashflow(ws):
    title(ws, "CashFlow", "Monthly cash flow, pre-financing, real RWF. Months 1–24 are the requested monthly view; the engine runs 120 months.")
    timeline(ws)
    c = mcol
    f, l = REF["CX_FIRST"], REF["CX_LAST"]
    pm, km = R("plant_m"), R("past_m")
    row = 8
    section(ws, row, "Sales (accrual)", 3 + MONTHS)
    row += 1
    sales_rows = []
    for key, cname, *_ in CROPS:
        rv, sz = ROW[f"{key}_REV"][1], ROW[f"{key}_SEASON"][1]
        month_rows(ws, f"CF_{key}", row, f"Fruit: {cname}", "RWF",
                   lambda m: f"=INDEX(Trees!$D${rv}:${YLAST}${rv},1,{c(m)}$4)*INDEX(Curves!$D${sz}:$O${sz},1,{c(m)}$5)")
        sales_rows.append(row)
        row += 1
    tr = ROW["TIM_REV"][1]
    month_rows(ws, "CF_TIM", row, "Boundary poles", "RWF", lambda m: f"=INDEX(Trees!$D${tr}:${YLAST}${tr},1,{c(m)}$4)/12")
    sales_rows.append(row)
    row += 1
    for code, species in SPECIES:
        month_rows(ws, f"CF_{code}_Y", row, f"{species}: young stock", "RWF", lambda m, k=code: f"={mcell(k + '_REV_Y', m)}")
        sales_rows.append(row)
        row += 1
        month_rows(ws, f"CF_{code}_C", row, f"{species}: culled adults", "RWF", lambda m, k=code: f"={mcell(k + '_REV_C', m)}")
        sales_rows.append(row)
        row += 1
    month_rows(ws, "CF_SALES", row, "Total sales", "RWF", lambda m: f"=SUM({c(m)}{sales_rows[0]}:{c(m)}{sales_rows[-1]})", bold=True, fill=FILL_TOTAL)
    sales = row
    row += 1
    month_rows(ws, "CF_KG", row, "Memo: fruit marketed", "kg",
               lambda m: "=" + "+".join(
                   f"INDEX(Trees!$D${ROW[k + '_PROD'][1]}:${YLAST}${ROW[k + '_PROD'][1]},1,{c(m)}$4)"
                   f"*INDEX(Curves!$D${ROW[k + '_SEASON'][1]}:$O${ROW[k + '_SEASON'][1]},1,{c(m)}$5)" for k, *_ in CROPS))
    memo_kg = row
    row += 2
    section(ws, row, "Cash in", 3 + MONTHS)
    row += 1
    month_rows(ws, "CF_COLLECT", row, "Collections from customers", "RWF",
               lambda m: (f"={c(m)}{sales}*(1-{R('CREDIT')})"
                          + (f"+{c(m - RECEIVABLE_LAG_M)}{sales}*{R('CREDIT')}" if m > RECEIVABLE_LAG_M else "")), bold=True)
    collect = row
    row += 2
    section(ws, row, "Operating costs", 3 + MONTHS)
    row += 1
    opex = [
        ("CF_PERM", "Permanent labour", lambda m: f"={R('WORKERS')}*{R('WAGE')}*{R('S_LABOR')}"),
        ("CF_CASUAL", "Casual labour (orchard)", lambda m: f"={R('L_Z2')}*{R('DAYS_HA')}*{R('DAYWAGE')}*{R('S_LABOR')}/12*IF({c(m)}$3>={pm},1,0)"),
        ("CF_VET", "Veterinary care, medicines, minerals", lambda m: f"=({mcell('GT_HEAD', m)}+{mcell('SH_HEAD', m)})*({R('VET')}+{R('MINERAL')})/12"),
        ("CF_FEED", "Feed bought", lambda m: f"={mcell('FD_BUYCOST', m)}"),
        ("CF_CONS", "Hay/silage conservation", lambda m: f"={mcell('FD_CONSCOST', m)}"),
        ("CF_FERT", "Orchard inputs", lambda m: f"={R('L_Z2')}*{R('FERT')}/12*IF({c(m)}$3>={pm},1,0)"),
        ("CF_PASTM", "Pasture and fodder upkeep", lambda m: f"=({R('L_PADDOCK')}+{R('L_FODDER')})*{R('PASTM')}/12*IF({c(m)}$3>={km},1,0)"),
        ("CF_REPL", "Replanting short-cycle crops", lambda m: f"=IF({c(m)}$3={pm}+12*({c(m)}$4-1),INDEX(Trees!$D${ROW['TR_REPL'][1]}:${YLAST}${ROW['TR_REPL'][1]},1,{c(m)}$4),0)"),
        ("CF_TRANSF", "Transport: fruit", lambda m: f"={c(m)}{memo_kg}*{R('TRANS_KG')}*{R('S_TRANS')}"),
        ("CF_TRANSL", "Transport: livestock", lambda m: f"=({mcell('GT_SOLDHD', m)}+{mcell('SH_SOLDHD', m)})*{R('TRANS_HD')}*{R('S_TRANS')}"),
        ("CF_PACK", "Packaging", lambda m: f"={c(m)}{memo_kg}*{R('PACK')}"),
        ("CF_MKTG", "Marketing and sales", lambda m: f"={c(m)}{sales}*{R('MKTG')}"),
        ("CF_COMMS", "Communications and software", lambda m: f"={R('COMMS')}"),
        ("CF_UTIL", "Utilities and fuel", lambda m: f"={R('UTIL')}"),
        ("CF_ADMIN", "Administration and accounting", lambda m: f"={R('ADMIN')}"),
        ("CF_INS", "Livestock insurance", lambda m: f"=({mcell('GT_VALUE', m)}+{mcell('SH_VALUE', m)})*{R('INS')}/12"),
        ("CF_RM", "Repairs and maintenance",
         lambda m: f'=SUMIF(CapEx!$I${f}:$I${l},"<"&{c(m)}$3,CapEx!$H${f}:$H${l})*(1+{R("CONT")})*{R("RM")}/12'),
        ("CF_LAND", "Land fees and taxes", lambda m: f"={R('LANDFEE')}/12"),
    ]
    op_first = row
    for key, label, fn in opex:
        month_rows(ws, key, row, label, "RWF", fn)
        row += 1
    op_last = row - 1
    month_rows(ws, "CF_OPEX", row, "Total operating costs", "RWF", lambda m: f"=SUM({c(m)}{op_first}:{c(m)}{op_last})", bold=True, fill=FILL_TOTAL)
    opex_row = row
    row += 2
    month_rows(ws, "CF_EBITDA", row, "EBITDA (sales less operating costs)", "RWF", lambda m: f"={c(m)}{sales}-{c(m)}{opex_row}", bold=True)
    ebitda = row
    row += 2
    section(ws, row, "Investment", 3 + MONTHS)
    row += 1
    month_rows(ws, "CF_CAPEX", row, "Capital items incl. contingency", "RWF",
               lambda m: f"=SUMIF(CapEx!$I${f}:$I${l},{c(m)}$3,CapEx!$H${f}:$H${l})*(1+{R('CONT')})")
    capex_row = row
    row += 1
    month_rows(ws, "CF_ANIMALS", row, "Animal purchases", "RWF",
               lambda m: f"={mcell('GT_BUYCOST', m)}+{mcell('SH_BUYCOST', m)}")
    animals_row = row
    row += 1
    month_rows(ws, "CF_INVEST", row, "Total investment", "RWF", lambda m: f"={c(m)}{capex_row}+{c(m)}{animals_row}", bold=True, fill=FILL_TOTAL)
    invest_row = row
    row += 2
    # Tax on the farm year's profit, paid in its last month. Computed here (not read from Annual)
    # so the monthly sheet stands alone; Metrics checks that it ties to Annual.
    month_rows(ws, "CF_TAX", row, "Income tax paid (year end)", "RWF",
               lambda m: (f"=IF(MOD({c(m)}$3,12)=0,MAX(0,SUMIF($D$4:${MLAST}$4,{c(m)}$4,$D${ebitda}:${MLAST}${ebitda})"
                          f"-{depreciation(c(m) + '$4')})*{R('TAX')},0)"))
    tax_row = row
    row += 2
    month_rows(ws, "CF_NET", row, "NET CASH FLOW (pre-financing)", "RWF",
               lambda m: f"={c(m)}{collect}-{c(m)}{opex_row}-{c(m)}{invest_row}-{c(m)}{tax_row}", bold=True, fill=FILL_TOTAL)
    net = row
    row += 1
    month_rows(ws, "CF_CUM", row, "Cumulative net cash flow", "RWF",
               lambda m: f"={(c(m - 1) + str(row)) if m > 1 else '0'}+{c(m)}{net}", bold=True)
    cum = row
    row += 1
    month_rows(ws, "CF_BAL", row, "Cash balance with founder capital", "RWF", lambda m: f"={R('capital')}+{c(m)}{cum}", bold=True)
    row += 1
    month_rows(ws, "CF_SHORT", row, "Shortfall flag (1 = capital exhausted)", "flag", lambda m: f"=IF({c(m)}{row - 1}<0,1,0)", "0")


def build_annual(ws):
    title(ws, "Annual", "Ten farm years: profit and loss (accrual) and cash flow (pre-financing). Real RWF.")
    year_header(ws, 3)
    yr = 3
    sumif = lambda key, y: f"=SUMIF(CashFlow!$D$4:${MLAST}$4,{ycol(y)}${yr},{mrange(key)})"  # noqa: E731
    r = 5
    section(ws, r, "Revenue", 3 + YEARS)
    r += 1
    rev_first = r
    lines = [(f"CF_{k}", f"Fruit: {n}") for k, n, *_ in CROPS] + [("CF_TIM", "Boundary poles")]
    for code, species in SPECIES:
        lines += [(f"CF_{code}_Y", f"{species}: young stock"), (f"CF_{code}_C", f"{species}: culled adults")]
    for key, label in lines:
        year_rows(ws, "AN_" + key, r, label, "RWF", lambda y, k=key: sumif(k, y))
        r += 1
    year_rows(ws, "AN_REV", r, "Total revenue", "RWF", lambda y: f"=SUM({ycol(y)}{rev_first}:{ycol(y)}{r - 1})", bold=True, fill=FILL_TOTAL)
    rev = r
    r += 2
    section(ws, r, "Direct costs", 3 + YEARS)
    r += 1
    d_first = r
    for key, label in [("CF_FEED", "Feed bought"), ("CF_CONS", "Hay/silage conservation"), ("CF_VET", "Veterinary care, medicines, minerals"),
                       ("CF_FERT", "Orchard inputs"), ("CF_PASTM", "Pasture and fodder upkeep"), ("CF_CASUAL", "Casual labour"),
                       ("CF_REPL", "Replanting short-cycle crops"), ("CF_TRANSF", "Transport: fruit"), ("CF_TRANSL", "Transport: livestock"),
                       ("CF_PACK", "Packaging")]:
        year_rows(ws, "AN_" + key, r, label, "RWF", lambda y, k=key: sumif(k, y))
        r += 1
    year_rows(ws, "AN_DIRECT", r, "Total direct costs", "RWF", lambda y: f"=SUM({ycol(y)}{d_first}:{ycol(y)}{r - 1})", bold=True, fill=FILL_TOTAL)
    direct = r
    r += 1
    year_rows(ws, "AN_GM", r, "Gross margin", "RWF", lambda y: f"={ycol(y)}{rev}-{ycol(y)}{direct}", bold=True)
    gm = r
    r += 1
    year_rows(ws, "AN_GMPCT", r, "Gross margin %", "share", lambda y: f"=IF({ycol(y)}{rev}>0,{ycol(y)}{gm}/{ycol(y)}{rev},0)", PCT)
    r += 2
    section(ws, r, "Overheads", 3 + YEARS)
    r += 1
    o_first = r
    for key, label in [("CF_PERM", "Permanent labour"), ("CF_MKTG", "Marketing and sales"), ("CF_COMMS", "Communications and software"),
                       ("CF_UTIL", "Utilities and fuel"), ("CF_ADMIN", "Administration and accounting"), ("CF_INS", "Livestock insurance"),
                       ("CF_RM", "Repairs and maintenance"), ("CF_LAND", "Land fees and taxes")]:
        year_rows(ws, "AN_" + key, r, label, "RWF", lambda y, k=key: sumif(k, y))
        r += 1
    year_rows(ws, "AN_OVER", r, "Total overheads", "RWF", lambda y: f"=SUM({ycol(y)}{o_first}:{ycol(y)}{r - 1})", bold=True, fill=FILL_TOTAL)
    over = r
    r += 2
    year_rows(ws, "AN_EBITDA", r, "EBITDA", "RWF", lambda y: f"={ycol(y)}{gm}-{ycol(y)}{over}", bold=True, fill=FILL_TOTAL)
    ebitda = r
    r += 1
    year_rows(ws, "AN_DEP", r, "Depreciation (incl. contingency)", "RWF", lambda y: "=" + depreciation(f"{ycol(y)}${yr}"))
    dep = r
    r += 1
    year_rows(ws, "AN_EBIT", r, "EBIT", "RWF", lambda y: f"={ycol(y)}{ebitda}-{ycol(y)}{dep}", bold=True)
    ebit = r
    r += 1
    year_rows(ws, "AN_TAX", r, "Income tax (no loss carry-forward modelled)", "RWF", lambda y: f"=MAX(0,{ycol(y)}{ebit})*{R('TAX')}")
    tax = r
    r += 1
    year_rows(ws, "AN_NP", r, "Net profit", "RWF", lambda y: f"={ycol(y)}{ebit}-{ycol(y)}{tax}", bold=True, fill=FILL_TOTAL)
    r += 2
    section(ws, r, "Cash flow (pre-financing)", 3 + YEARS)
    r += 1
    year_rows(ws, "AN_COLLECT", r, "Collections", "RWF", lambda y: sumif("CF_COLLECT", y))
    col_ = r
    year_rows(ws, "AN_OPEX", r + 1, "Operating costs", "RWF", lambda y: sumif("CF_OPEX", y))
    year_rows(ws, "AN_CAPEX", r + 2, "Capital items incl. contingency", "RWF", lambda y: sumif("CF_CAPEX", y))
    year_rows(ws, "AN_ANIMALS", r + 3, "Animal purchases", "RWF", lambda y: sumif("CF_ANIMALS", y))
    year_rows(ws, "AN_TAXPAID", r + 4, "Income tax paid", "RWF", lambda y: sumif("CF_TAX", y))
    year_rows(ws, "AN_NET", r + 5, "Net cash flow", "RWF",
              lambda y: f"={ycol(y)}{col_}-{ycol(y)}{col_ + 1}-{ycol(y)}{col_ + 2}-{ycol(y)}{col_ + 3}-{ycol(y)}{col_ + 4}", bold=True, fill=FILL_TOTAL)
    net = r + 5
    year_rows(ws, "AN_CUM", r + 6, "Cumulative net cash flow", "RWF",
              lambda y: f"={(ycol(y - 1) + str(r + 6)) if y > 1 else '0'}+{ycol(y)}{net}", bold=True)
    cum = r + 6
    year_rows(ws, "AN_DF", r + 7, "Discount factor (end of year)", "factor", lambda y: f"=1/(1+{R('DR')})^{ycol(y)}${yr}", NUM2)
    year_rows(ws, "AN_PV", r + 8, "Present value of net cash flow", "RWF", lambda y: f"={ycol(y)}{net}*{ycol(y)}{r + 7}")
    year_rows(ws, "AN_PAYBACK", r + 9, "Payback marker", "year",
              lambda y: (f'=IF(AND({ycol(y)}{cum}>=0,{ycol(y - 1)}{cum}<0),{ycol(y)}${yr},"")' if y > 1
                         else f'=IF({ycol(y)}{cum}>=0,{ycol(y)}${yr},"")'), YEARFMT)
    year_rows(ws, "AN_BE", r + 10, "Operating break-even marker (EBITDA > 0)", "year",
              lambda y: f'=IF({ycol(y)}{ebitda}>0,{ycol(y)}${yr},"")', YEARFMT)
    r += 12
    section(ws, r, "Operating indicators", 3 + YEARS)
    r += 1
    year_rows(ws, "AN_KG", r, "Fruit marketed", "kg", lambda y: sumif("CF_KG", y))
    for i, (code, species) in enumerate(SPECIES, start=1):
        fc = ROW[f"{code}_FCLOSE"][1]
        year_rows(ws, f"AN_{code}_F", r + i, f"{species}: breeding females at year end", "head",
                  lambda y, rr=fc: f"=INDEX(Livestock!$D${rr}:${MLAST}${rr},1,12*{ycol(y)}${yr})", NUM1)
    year_rows(ws, "AN_FDEM", r + 3, "Feed demand", "kg DM", lambda y: sumif("FD_DEMAND", y))
    year_rows(ws, "AN_FBUY", r + 4, "Feed bought", "kg DM", lambda y: sumif("FD_BUY", y))
    year_rows(ws, "AN_FSELF", r + 5, "Feed self-sufficiency", "share",
              lambda y: f"=IF({ycol(y)}{r + 3}>0,1-{ycol(y)}{r + 4}/{ycol(y)}{r + 3},1)", PCT)
    ws.freeze_panes = "D4"


def build_metrics(ws):
    title(ws, "Metrics", "Headline results for the selected scenario and model. Placeholder inputs make these illustrative only.")
    header_row(ws, 4, ["Key", "Metric", "Value", "Unit / note"])
    an = lambda key: yrange(key)  # noqa: E731
    cf_cum = mrange("CF_CUM")
    rows = [
        ("MX_SCEN", "Scenario", f"={R('scenario_name')}", "", None),
        ("MX_MODEL", "Farm model", f"={R('model_name')}", "", None),
        ("MX_CAP", "Capital available, years 1–2", f"={R('capital')}", "RWF", NUM),
        ("MX_INV10", "Total investment over 10 years (capital items + animals)", f"=SUM({an('AN_CAPEX')})+SUM({an('AN_ANIMALS')})", "RWF", NUM),
        ("MX_INV24", "Investment in months 1–24", f"=SUM(CashFlow!$D${ROW['CF_INVEST'][1]}:${mcol(24)}${ROW['CF_INVEST'][1]})", "RWF", NUM),
        ("MX_PEAK", "Peak funding need (lowest cumulative cash)", f"=MAX(0,-MIN({cf_cum}))", "RWF", NUM),
        ("MX_PEAKM", "Month of peak funding need", f"=MATCH(MIN({cf_cum}),{cf_cum},0)", "model month", MONTHFMT),
        ("MX_GAP", "FUNDING GAP (peak need less capital available)", f"=MAX(0,C10-C7)", "RWF", NUM),
        ("MX_GAPUSD", "Funding gap in USD", f"=C12/{R('FX')}", "USD at the Assumptions rate", NUM),
        ("MX_NPV", "NPV, 10 years, real, no terminal value", f"=NPV({R('DR')},{an('AN_NET')})", "RWF", NUM),
        ("MX_IRR", "IRR, 10 years", f'=IFERROR(IRR({an("AN_NET")}),"n/a")', "share", PCT),
        ("MX_PAYBACK", "Payback year (cumulative cash turns positive)", f'=IF(COUNT({an("AN_PAYBACK")})=0,"beyond Y10",MIN({an("AN_PAYBACK")}))', "farm year", YEARFMT),
        ("MX_BE", "First year with positive EBITDA", f'=IF(COUNT({an("AN_BE")})=0,"beyond Y10",MIN({an("AN_BE")}))', "farm year", YEARFMT),
        ("MX_REV5", "Revenue, year 5", f"=Annual!{ycol(5)}{ROW['AN_REV'][1]}", "RWF", NUM),
        ("MX_REV10", "Revenue, year 10", f"=Annual!{ycol(10)}{ROW['AN_REV'][1]}", "RWF", NUM),
        ("MX_EBITDA5", "EBITDA, year 5", f"=Annual!{ycol(5)}{ROW['AN_EBITDA'][1]}", "RWF", NUM),
        ("MX_EBITDA10", "EBITDA, year 10", f"=Annual!{ycol(10)}{ROW['AN_EBITDA'][1]}", "RWF", NUM),
        ("MX_ROIC", "Return on invested capital, average Y6–Y10 (after tax)",
         f'=IFERROR(AVERAGE(Annual!{ycol(6)}{ROW["AN_EBIT"][1]}:{ycol(10)}{ROW["AN_EBIT"][1]})*(1-{R("TAX")})/C8,"n/a")', "share", PCT),
        ("MX_FSELF5", "Feed self-sufficiency, year 5", f"=Annual!{ycol(5)}{ROW['AN_FSELF'][1]}", "share", PCT),
        ("MX_TIMBER", "Memo: standing boundary-timber value at Y10 (not in NPV)", f"={R('TIM_MEMO')}", "RWF", NUM),
    ]
    r = 5
    for key, label, formula, unit, fmt in rows:
        put(ws, f"A{r}", key, F_NOTE)
        bold = key in ("MX_GAP", "MX_NPV", "MX_PEAK")
        put(ws, f"B{r}", label, F_BOLD if bold else F_BASE)
        fx(ws, f"C{r}", formula, fmt or "General", bold=bold, fill=FILL_KEY if key == "MX_GAP" else None)
        put(ws, f"D{r}", unit, F_NOTE)
        REF[key] = f"Metrics!$C${r}"
        r += 1
    # The funding-gap formula above refers to C10 (peak) and C7 (capital) by position; assert that holds.
    assert REF["MX_PEAK"] == "Metrics!$C$10" and REF["MX_CAP"] == "Metrics!$C$7" and REF["MX_GAP"] == "Metrics!$C$12"
    assert REF["MX_INV10"] == "Metrics!$C$8"
    r += 1
    section(ws, r, "Checks (all should read OK)", 4)
    r += 1
    sales_total = f"SUM({mrange('CF_SALES')})"
    checks = [
        ("Land shares sum to 100%", f'=IF(ABS({R("CHK_LAND")}-1)<0.0001,"OK","CHECK")'),
        ("Orchard mix sums to 100%", f'=IF(ABS({R("CHK_ORCH")}-1)<0.0001,"OK","CHECK")'),
        ("Monthly sales tie to annual revenue", f'=IF(ABS({sales_total}-SUM({yrange("AN_REV")}))<1,"OK","CHECK")'),
        ("Monthly net cash ties to annual net cash", f'=IF(ABS(SUM({mrange("CF_NET")})-SUM({yrange("AN_NET")}))<1,"OK","CHECK")'),
        ("Monthly tax ties to annual tax", f'=IF(ABS(SUM({mrange("CF_TAX")})-SUM({yrange("AN_TAX")}))<1,"OK","CHECK")'),
        ("Herd never negative", f'=IF(MIN(MIN({mrange("GT_FCLOSE")}),MIN({mrange("SH_FCLOSE")}))>=-0.0001,"OK","CHECK")'),
        ("Harvest seasonality sums to 1 for every crop",
         "=IF(AND(" + ",".join(f"ABS(Curves!$P${ROW[k + '_SEASON'][1]}-1)<0.001" for k, *_ in CROPS) + '),"OK","CHECK")'),
        ("Pasture growth averages 1", f'=IF(ABS(Curves!$P${ROW["PGROWTH"][1]}-1)<0.001,"OK","CHECK")'),
    ]
    for label, formula in checks:
        put(ws, f"B{r}", label)
        fx(ws, f"C{r}", formula, "General")
        r += 1
    widths(ws, {"A": 12, "B": 58, "C": 18, "D": 28})


def main():
    wb = Workbook()
    ws_cover = wb.active
    ws_cover.title = "Cover"
    sheets = {name: wb.create_sheet(name) for name in
              ["Control", "Assumptions", "Models", "Curves", "Sensitivity", "Land", "Trees", "Livestock",
               "Feed", "CapEx", "CashFlow", "Annual", "Metrics"]}
    build_cover(ws_cover)
    build_control(sheets["Control"])
    build_assumptions(sheets["Assumptions"])
    sheets["Control"][REF["_capital_usd_cell"]] = f"={R('capital')}/{R('FX')}"
    sheets["Control"][REF["_capital_usd_cell"]].number_format = NUM
    sheets["Control"][REF["_capital_usd_cell"]].font = F_LINK
    build_models(sheets["Models"])
    build_curves(sheets["Curves"])
    build_sensitivity(sheets["Sensitivity"])
    build_land(sheets["Land"])
    build_trees(sheets["Trees"])
    build_livestock(sheets["Livestock"])
    build_feed(sheets["Feed"])
    build_capex(sheets["CapEx"])
    build_cashflow(sheets["CashFlow"])
    build_annual(sheets["Annual"])
    build_metrics(sheets["Metrics"])
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
