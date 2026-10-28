"""Write every output of the studio package from the model: SVG drawings, GLB + parts.json, DXF
cutting sheets, CSV/XLSX tables. Deterministic: no timestamps, sorted ids, fixed formatting."""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import zipfile
from pathlib import Path

import params as P
import sheets as S
from draw import FILL, Sheet, iso_draw
from joinery import HARDWARE, bag
from nest import nest
from out_glb import write as write_glb
from prices import CONTINGENCY, LABOUR, PRICES, SHARED
from services import FIT_SIZE, GROUP_COLOUR, SIZE, SYSTEM_COLOUR, SYSTEM_GROUP

SECTION_KG = {"RHS 150x100x5": 18.2, "RHS 80x40x3": 5.2, "SHS 40x40x3": 3.4, "SHS 80x80x4": 9.2, "SHS 60x60x3": 5.2,
              "ledger": 5.2}


def wcsv(path: Path, header: list[str], rows: list[list]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow([f"{v:.2f}".rstrip("0").rstrip(".") if isinstance(v, float) else v for v in r])


def edge_m(p) -> float:
    code = p.edges
    if not code:
        return 0.0
    n = 0.0
    if "L2" in code:
        n += 2 * p.length
    elif "L1" in code:
        n += p.length
    if "S2" in code:
        n += 2 * p.width
    elif "S1" in code:
        n += p.width
    return n / 1000


def steel_kg(p) -> float:
    L = max(p.box.size) / 1000
    for k, v in SECTION_KG.items():
        if k in p.name or (k == "ledger" and "ledger" in p.name):
            return v * L
    if "tread" in p.name:
        return 12.0
    if "landing" in p.name:
        return 60.0
    return 6.0 * L


def all_parts(m):
    return [p for k in m["kits"] for p in k.parts]


# ------------------------------------------------------------------------------- tables
def bom(m, layouts) -> tuple[list[list], dict]:
    rows: list[list] = []

    def add(kit, cat, item, qty, unit, key, trade=""):
        if not qty:
            return
        price, label, src, note = PRICES.get(key, (0, "U", "", "no price"))
        rows.append([kit, cat, trade, item, round(qty, 2), unit, price, label, src, round(qty * price), note])

    counts: dict[str, int] = {}
    for sl in layouts:
        counts[sl.material] = counts.get(sl.material, 0) + 1
    for mat, n in sorted(counts.items()):
        add("K1-K7", "boards", f"{mat} sheet {P.SHEETS[mat][0]}x{P.SHEETS[mat][1]}x{P.SHEETS[mat][2]}", n, "sheet",
            f"sheet_{mat}", "carpentry")
    eb = sum(edge_m(p) for p in all_parts(m))
    add("K1-K7", "boards", "ABS edge band (2 mm visible, 0.8 mm hidden)", round(eb * 1.1, 1), "m", "edgeband", "carpentry")
    furn = [p for p in all_parts(m) if p.kit in ("K2", "K3", "K4", "K5", "K6", "K7")]
    hw = bag(furn)
    for k in m["kits"]:
        for st in k.steps:
            for code, n in st.hardware.items():
                if code in ("F", "L", "C", "K"):
                    hw[code] = hw.get(code, 0) + n
    for code in sorted(hw):
        name, key = HARDWARE[code]
        add("K2-K7", "hardware", f"Bag {code}: {name}", hw[code] * 1.05, "each", key, "carpentry")
    roof_steel = sum(steel_kg(p) for p in all_parts(m) if p.material == "STEEL")
    roof_timber = sum(p.box.size[0] * p.box.size[1] * p.box.size[2] / 1e9 * (0.35 if "chair" in p.name.lower() or "table" in p.name.lower() else 1)
                      for p in all_parts(m) if p.material == "TIMBER")
    roof_poly = sum(p.box.size[0] * p.box.size[1] / 1e6 for p in all_parts(m) if p.material == "POLY")
    roof_cable = sum(max(p.box.size[0], p.box.size[1]) / 1000 * 10 for p in all_parts(m) if p.material == "CABLE")
    add("K11", "roof kit", "Fabricated steel: beams, joists, posts, rails, treads (painted)", roof_steel, "kg", "steel_kg", "welding")
    add("K11", "roof kit", "Treated hardwood: deck boards, top rail, chairs, side table", roof_timber, "m3", "timber_m3",
        "carpentry")
    add("K11", "roof kit", "Tinted multiwall polycarbonate 10 mm (front deck canopy)", roof_poly * 1.1, "m2", "poly_m2",
        "welding")
    add("K11", "roof kit", "Guardrail cable infill: 4 mm stainless, 10 runs, with tensioners", roof_cable, "m", "cable_m",
        "welding")
    add("K11", "roof kit", "Stair stringers, handrails and footings", 1, "set", "stairframe", "welding")
    for k in m["kits"]:
        for b in k.bought:
            if b.category in ("existing",) or not b.qty or not b.price:
                continue
            if b.price in ("batten", "stud") and b.unit == "each":
                continue
            add(k.id, b.category, b.name, b.qty, b.unit, b.price, b.trade)
    nb = sum(1 for k in m["kits"] for b in k.bought if b.price == "batten")
    ns = sum(1 for k in m["kits"] for b in k.bought if b.price == "stud")
    add("K1", "timber", "Battens 38 x 50 treated, 2.4 m", nb * 1.1, "each", "batten", "carpentry")
    add("K2", "timber", "Studs 40 x 40 treated, 2.4 m", ns * 1.1, "each", "stud", "carpentry")
    for name, rwf, label in LABOUR:
        rows.append(["labour", "labour", "", name, 1, "lot", rwf, label, "AS-21", rwf, ""])
    unit = sum(r[9] for r in rows)
    shared = sum(r for _, r, _ in SHARED)
    verified = sum(r[9] for r in rows if r[7] == "V")
    tot = {"unit_materials_labour": unit, "unit_with_contingency": unit * (1 + CONTINGENCY), "shared_site": shared,
           "shared_with_contingency": shared * (1 + CONTINGENCY),
           "units": P.UNITS, "all_units_total": (P.UNITS * unit + shared) * (1 + CONTINGENCY), "verified_share": verified / unit if unit else 0,
           "usd_rate": 1475.91}
    return rows, tot


def write_tables(m, out: Path, layouts) -> dict:
    d = out / "data"
    parts = sorted(all_parts(m), key=lambda p: p.id)
    wcsv(d / "cut-list.csv", ["id", "kit", "name", "material", "length_mm", "width_mm", "thickness_mm", "edges",
                              "edge_band_m", "face_ops", "edge_ops", "step"],
         [[p.id, p.kit, p.name, p.material, round(p.length), round(p.width), round(p.t), p.edges, round(edge_m(p), 2),
           sum(1 for o in p.ops if o.cnc), sum(1 for o in p.ops if not o.cnc), p.step]
          for p in parts if p.sheet and p.material in P.SHEETS])
    wcsv(d / "steel-list.csv", ["id", "name", "material", "length_mm", "kg_or_m3", "step"],
         [[p.id, p.name, p.material, round(max(p.box.size)),
           round(steel_kg(p), 1) if p.material == "STEEL" else round(p.box.size[0] * p.box.size[1] * p.box.size[2] / 1e9, 3),
           p.step] for p in parts if not p.sheet])
    rows, tot = bom(m, layouts)
    wcsv(d / "bom.csv", ["kit", "category", "trade", "item", "qty_per_unit", "unit", "unit_price_rwf", "label", "source",
                         "total_per_unit_rwf", "note"], rows)
    cats: dict[str, float] = {}
    for r in rows:
        cats[r[1]] = cats.get(r[1], 0) + r[9]
    summary = [[c, round(v), round(P.UNITS * v)] for c, v in sorted(cats.items(), key=lambda t: -t[1])]
    summary += [["shared site works (all units)", "", round(tot["shared_site"])],
                [f"contingency {CONTINGENCY * 100:.0f} %", round(tot["unit_materials_labour"] * CONTINGENCY),
                 round((P.UNITS * tot["unit_materials_labour"] + tot["shared_site"]) * CONTINGENCY)],
                ["TOTAL", round(tot["unit_with_contingency"]), round(tot["all_units_total"])]]
    wcsv(d / "bom-summary.csv", ["category", "per_unit_rwf", "all_units_rwf"], summary)
    wcsv(d / "shared-works.csv", ["item", "rwf", "label"], [[n, round(v), lab] for n, v, lab in SHARED])
    hw = [[code, HARDWARE[code][0]] for code in sorted(HARDWARE)]
    wcsv(d / "hardware.csv", ["bag", "item"], hw)
    wcsv(d / "checks.csv", ["id", "rule", "result", "detail"], [[r["id"], r["rule"], r["result"], r["detail"]] for r in m["checks"]])
    wcsv(d / "circuits.csv", ["circuit", "serves", "device", "cable_mm2", "design_A", "derated_Iz_A", "run_m", "drop_pct",
                              "points", "group", "colour", "relay"],
         [[c.id, c.name, c.breaker, c.cable, round(c.Ib, 1) if c.loads else "", round(c.Iz, 1),
           c.length_m if c.loads else "", round(c.vdrop_pct(), 2) if c.loads else "", " ".join(c.loads), c.group,
           GROUP_COLOUR[c.group], "pauses while the hob runs" if c.relay else ""] for c in m["circuits"]])
    wcsv(d / "points.csv", ["point", "kind", "circuit", "description", "x_mm", "y_mm", "z_mm", "watts", "face", "model",
                            "w_mm", "h_mm", "d_mm"],
         [[p.id, p.kind, p.circuit, p.label, round(p.x), round(p.y), round(p.z), round(p.watts), p.face, p.model,
           *SIZE[p.model]] for p in m["points"]])
    wcsv(d / "demand.csv", ["item", "amps"], [[k, round(v, 1)] for k, v in m["md"].items()] +
         [["inlet_A", m["inlet"]["rating_A"]], ["main_A", m["inlet"]["main_A"]], ["tails_mm2", m["inlet"]["tails_mm2"]]])
    wcsv(d / "pipes.csv", ["run", "system", "dn", "length_m", "fall", "note", "colour"],
         [[r.id, r.system, r.dn, round(r.length / 1000, 1), f"1:{r.fall():.0f}" if r.system in ("soil", "waste") else "",
           r.note, SYSTEM_COLOUR[SYSTEM_GROUP[r.system]]] for r in m["runs"]])
    wcsv(d / "fittings.csv", ["point", "kind", "model", "runs", "description", "item", "x_mm", "y_mm", "z_mm", "face",
                              "rating", "w_mm", "h_mm", "d_mm"],
         [[f.id, f.kind, f.model, f.runs, f.label, f.item, round(f.x), round(f.y), round(f.z), f.face, f.rating,
           *FIT_SIZE[f.model]] for f in m["fittings"]])
    wcsv(d / "site-water.csv", ["item", "value"], [[k, v] for k, v in m["site"].items()])
    wcsv(d / "steps.csv", ["step", "kit", "title", "parts", "hardware", "people", "note"],
         [[s.id, s.kit, s.title, len(s.adds), " ".join(f"{k}x{v}" for k, v in s.hardware.items() if v), s.persons, s.note]
          for s in sorted((s for k in m["kits"] for s in k.steps), key=lambda s: s.id)])
    wcsv(d / "nesting.csv", ["material", "sheet", "parts", "utilisation_pct"],
         [[sl.material, sl.index, len(sl.placed), round(sl.utilisation * 100, 1)] for sl in layouts])
    wcsv(d / "kits.csv", ["kit", "name", "trade", "made_where", "parts", "bought_lines"],
         [[k.id, k.name, k.trade, k.where, len(k.parts), sum(1 for b in k.bought if b.qty)] for k in m["kits"] if k.id != "K0"])
    write_xlsx(d / "bom.xlsx", rows, summary, tot)
    tot["sheets"] = {mat: sum(1 for sl in layouts if sl.material == mat) for mat in sorted({sl.material for sl in layouts})}
    (d / "totals.json").write_text(json.dumps(tot, indent=1, sort_keys=True) + "\n")
    return tot


def write_xlsx(path: Path, rows, summary, tot) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    fixed = dt.datetime(2026, 10, 26)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = "farm scripts/studio"
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Container studios - bill of materials", f"rev {P.REV} {P.REV_DATE}", P.STATUS])
    ws.append(["Prices labelled A are placeholders (AS-21): three quotes per line before buying."])
    ws.append([])
    ws.append(["category", "per unit RWF", f"{P.UNITS} units RWF"])
    for r in summary:
        ws.append(r)
    ws.append([])
    ws.append(["USD at BNR 1,475.91 (2026-10-23, STU-36)", "", round(tot["all_units_total"] / 1475.91)])
    ws2 = wb.create_sheet("BOM per unit")
    ws2.append(["kit", "category", "trade", "item", "qty", "unit", "unit price", "label", "source", "total", "note"])
    for r in rows:
        ws2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name in sorted(src.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, src.read(name))


# ------------------------------------------------------------------------------- DXF
def write_dxf(out: Path, layouts) -> None:
    import ezdxf
    ezdxf.options.write_fixed_meta_data_for_testing = True
    for sl in layouts:
        doc = ezdxf.new("R2000")
        doc.units = 4  # mm
        msp = doc.modelspace()
        msp.add_lwpolyline([(0, 0), (sl.L, 0), (sl.L, sl.W), (0, sl.W)], close=True, dxfattribs={"layer": "SHEET"})
        for pl in sl.placed:
            p = pl.part
            x0, y0 = pl.x, pl.y
            msp.add_lwpolyline([(x0, y0), (x0 + pl.l, y0), (x0 + pl.l, y0 + pl.w), (x0, y0 + pl.w)], close=True,
                               dxfattribs={"layer": "CUT"})
            msp.add_text(p.id, height=min(25, pl.w / 4), dxfattribs={"layer": "LABEL"}).set_placement(
                (x0 + pl.l / 2, y0 + pl.w / 2), align=ezdxf.enums.TextEntityAlignment.MIDDLE_CENTER)
            for o in p.ops:
                if not o.cnc:
                    continue
                u, v = (o.v, o.u) if not pl.swapped else (o.u, o.v)
                if p.w >= p.h:  # first in-plane axis is the long one -> along the sheet
                    u, v = o.u, o.v
                else:
                    u, v = o.v, o.u
                layer = o.layer + ("_FACEB" if o.face == "B" else "")
                if o.kind == "cutout":
                    msp.add_lwpolyline([(x0 + u, y0 + v), (x0 + u + (o.w if p.w >= p.h else o.h), y0 + v),
                                        (x0 + u + (o.w if p.w >= p.h else o.h), y0 + v + (o.h if p.w >= p.h else o.w)),
                                        (x0 + u, y0 + v + (o.h if p.w >= p.h else o.w))], close=True,
                                       dxfattribs={"layer": layer})
                else:
                    msp.add_circle((x0 + u, y0 + v), o.dia / 2, dxfattribs={"layer": layer})
        doc.saveas(out / "cnc" / f"{sl.material}-sheet-{sl.index:02d}.dxf")


# ------------------------------------------------------------------------------- model
def write_model(m, out: Path) -> None:
    kitname = {k.id: k.name for k in m["kits"]}
    steps = {s.id: s.title for k in m["kits"] for s in k.steps}
    nodes, meta = [], []
    for k in m["kits"]:
        for p in k.parts:
            b = p.box
            nodes.append({"id": p.id, "mat": p.material, "box": (b.x0, b.y0, b.z0, b.x1, b.y1, b.z1),
                          "extras": {"kit": k.id, "layer": p.layer, "step": int(p.step[1:]) if p.step else 0,
                                     "explode": [round(p.explode[0] / 1000, 3), round(p.explode[2] / 1000, 3),
                                                 round(p.explode[1] / 1000, 3)]}})
            meta.append({"id": p.id, "name": p.name, "kit": k.id, "kitName": kitname[k.id], "layer": p.layer,
                         "material": p.material, "size": [round(p.length), round(p.width), round(p.t)],
                         "step": p.step, "stepTitle": steps.get(p.step, "")})
        for bt in k.bought:
            if bt.box is None:
                continue
            b = bt.box
            nodes.append({"id": bt.id, "mat": bt.color or "white", "box": (b.x0, b.y0, b.z0, b.x1, b.y1, b.z1),
                          "extras": {"kit": k.id, "layer": bt.layer, "step": int(bt.step[1:]) if bt.step else 0,
                                     "explode": [0, 0, 0]}})
            meta.append({"id": bt.id, "name": bt.name, "kit": k.id, "kitName": kitname[k.id], "layer": bt.layer,
                         "material": bt.color or "bought", "size": [round(v) for v in b.size], "step": bt.step,
                         "stepTitle": steps.get(bt.step, "")})
    nodes.sort(key=lambda n: n["id"])
    meta.sort(key=lambda n: n["id"])
    write_glb(out / "model" / "studio.glb", nodes)
    layers = sorted({n["layer"] for n in meta})
    (out / "model" / "parts.json").write_text(json.dumps({"rev": P.REV, "status": P.STATUS, "layers": layers,
                                                          "steps": dict(sorted(steps.items())), "parts": meta},
                                                         separators=(",", ":"), sort_keys=True) + "\n")
    # poster: the furnished unit without the front wall and roof
    sh = Sheet("", "")
    lo = m["lo"]
    its = []
    for (b, fl, nm, i, k, layer) in S.items(m, layers=["furniture", "fixtures", "envelope", "shell"], exclude_layers=()):
        if i.startswith(("K0-FRONT", "K0-ROOF", "K0-SLIDER", "K0-ENDB", "K1-BAT", "K1-INS", "K1-F", "K2-CB-ENDB", "K2-CB-FRONT")) or b.z0 >= lo.H - 1:
            continue
        if (layer == "envelope" and b.y0 >= lo.W - 1) or (i.startswith("K7-") and b.y1 >= lo.W - 1):
            continue
        its.append((b, fl, "normal", ""))
    iso_draw(sh, its, 210, 150, (380, 250), max_scale=0.07)
    svg = sh.svg()
    svg = svg[:svg.index("<rect x=\"10\"")] + "".join(sh.body) + "</svg>\n"
    (out / "model" / "studio-poster.svg").write_text(svg)


# ------------------------------------------------------------------------------- drawings
def write_drawings(m, out: Path, layouts) -> list[tuple[str, str, str]]:
    lo = m["lo"]
    sheets: list[Sheet] = [S.a101_existing(m), S.a102_options(m), S.a103_plan(m), S.a201_elev(m), S.a202_ends(m),
                           S.a401_details(m), S.s101_roof(m), S.e101_layout(m), S.e201_sld(m), S.e202_circuits(m),
                           S.p101_plan(m), S.p201_iso(m), S.p301_site(m)]
    for i, kid in enumerate(("K1", "K2", "K3", "K4", "K5", "K6", "K7", "K11")):
        sheets.append(S.j_kit(m, kid, f"J-{101 + i}"))
    for i in range(0, len(layouts), 4):
        sheets.append(S.c_cutting(layouts, i, f"C-{101 + i // 4}"))
    index = [(s.number, s.title) for s in sheets]
    cover = S.a001_cover(m, [("A-001", "Cover and drawing index")] + index)
    sheets.insert(0, cover)
    steps = sorted(((s, k) for k in m["kits"] for s in k.steps), key=lambda t: t[0].id)
    done: set[str] = set()
    for st, k in steps:
        sheets.append(S.step_sheet(m, st, k, done))
        done |= set(st.adds)
    sheets.append(S.template_sheet())
    files = []
    for s in sheets:
        fn = f"{s.number.lower()}.svg"
        (out / "drawings" / fn).write_text(s.svg())
        files.append((s.number, s.title, fn))
    wcsv(out / "data" / "drawings.csv", ["number", "title", "file"], [list(f) for f in files])
    (out / "render" / "plan.svg").write_text(S.plan_bw(m))
    (out / "render" / "scene.json").write_text(json.dumps(S.scene_json(m), indent=1) + "\n")
    _ = lo
    return files


def write_dims(m, out: Path) -> None:
    """The measurements the trade guides quote (guide/*.md {{studio.*}} tokens), in mm unless named."""
    from layout import AISLE_MIN, GALLEY, GALLEY_D, WORKTOP_D
    lo = m["lo"]
    k1 = next(k for k in m["kits"] if k.id == "K1")
    bought = {b.id: b for b in k1.bought}
    bat = [b for b in k1.bought if b.price == "batten"]
    foam = sum(b.qty for b in k1.bought if b.price == "sprayfoam")
    r = lambda v: int(round(v))
    d = {
        "rev": P.REV, "rev_date": P.REV_DATE, "status": P.STATUS, "units": P.UNITS,
        "int_l": r(P.m("int_length", P.INT_L)), "int_w": r(P.m("int_width", P.INT_W)), "int_h": r(P.m("int_height", P.INT_H)),
        "l": r(lo.L), "w": r(lo.W), "h": r(lo.H), "area_m2": round(lo.L * lo.W / 1e6, 1),
        "lining_wall": r(P.LINING_WALL.value), "lining_ceil": r(P.LINING_CEIL.value), "lining_floor": r(P.LINING_FLOOR.value),
        "floor_sheets": r(bought["K1-FLOOR"].qty), "floor_m2": round((lo.L + 2 * P.LINING_WALL.value) * (lo.W + 2 * P.LINING_WALL.value) / 1e6, 1),
        "floor_screws": next(s.hardware.get("F", 0) for s in k1.steps if s.id == "S02"),
        "membrane_m2": bought["K1-DPM"].qty, "plank_m2": bought["K1-SPC"].qty,
        "battens": len(bat), "batten_screws": next(s.hardware.get("L", 0) for s in k1.steps if s.id == "S03"),
        "lining_screws": next(s.hardware.get("F", 0) for s in k1.steps if s.id == "S05"),
        "foam_m2": round(foam, 1), "lining_panels": sum(1 for p in k1.parts),
        "slider_x0": r(lo.o1[0]), "slider_x1": r(lo.o1[1]), "slider_head": r(lo.o1_head),
        "w1_x0": r(lo.o2[0]), "w1_x1": r(lo.o2[1]), "w1_sill": r(lo.w1_sill), "w1_head": r(lo.w1_head),
        "bed_x1": r(lo.bed_x1), "bed_y0": r(lo.bed_y0), "bed_y1": r(lo.bed_y1),
        "bedside_x0": r(lo.bedside[0]), "bedside_x1": r(lo.bedside[1]),
        "sofa_x0": r(lo.sofa[0]), "sofa_x1": r(lo.sofa[1]),
        "galley_x0": r(lo.galley_x0), "galley_d": GALLEY_D, "worktop_d": WORKTOP_D, "galley_run": sum(GALLEY),
        "fridge_w": GALLEY[0], "hob_w": GALLEY[1], "sink_w": GALLEY[2],
        "aisle": r(lo.part_x0 - WORKTOP_D - lo.sofa[1]), "aisle_min": AISLE_MIN,
        "part_x0": r(lo.part_x0), "part_t": r(lo.part_t), "bath_x0": r(lo.bath_x0), "wall_b": r(lo.wall_b),
        "tray_l": r(lo.tray[1] - lo.tray[0]), "tray_w": r(lo.tray[3] - lo.tray[2]),
        "heater_l": r(P.HEATER_L.value),
        "parts": sum(1 for k in m["kits"] for p in k.parts if p.sheet and p.material in P.SHEETS),
        "sheets": {mat: sum(1 for sl in m["layouts"] if sl.material == mat) for mat in sorted({sl.material for sl in m["layouts"]})},
    }
    d.update({f"sheets_{mat.lower()}": n for mat, n in d["sheets"].items()})
    tot = json.loads((out / "data" / "totals.json").read_text())  # written by write_tables, just before
    d.update({"unit_rwf_m": round(tot["unit_with_contingency"] / 1e6, 1), "all_units_rwf_m": round(tot["all_units_total"] / 1e6, 1),
              "all_units_usd_k": round(tot["all_units_total"] / tot["usd_rate"] / 1000), "shared_rwf_m": round(tot["shared_site"] / 1e6, 1),
              "persons": r(m["site"]["persons"]), "daily_l": r(m["site"]["daily_L"]), "tank_m3": m["site"]["tank_m3"],
              "septic_m3": m["site"]["septic_m3"], "septic_c1_m3": m["site"]["chamber1_m3"], "septic_c2_m3": m["site"]["chamber2_m3"],
              "septic_pop": r(P.SEPTIC_POP.value), "unit_ml_rwf_m": round(tot["unit_materials_labour"] / 1e6, 1)})
    cats = {row["category"]: row for row in csv.DictReader((out / "data" / "bom-summary.csv").open(encoding="utf-8"))}
    d["roof_rwf_m"] = round(float(cats["roof kit"]["per_unit_rwf"]) / 1e6, 1)
    (out / "data" / "dims.json").write_text(json.dumps(d, indent=1, sort_keys=True) + "\n")
    # batten set-out: each batten's centreline from the bare steel, as the carpenter marks it before any lining
    e = P.LINING_WALL.value
    rows = []
    for b in bat:
        where = b.name.split("(")[1].rstrip(")")
        bx = b.box
        if where == "end A":
            rows.append([b.id, where, "from the steel of the back wall", r((bx.y0 + bx.y1) / 2 + e), r(bx.z1 - bx.z0)])
        elif where == "slider head":
            rows.append([b.id, where, "horizontal, from the steel of end A", f"{r(bx.x0 + e)}-{r(bx.x1 + e)}",
                         r(bx.x1 - bx.x0)])
        else:
            rows.append([b.id, where, "from the steel of end A", r((bx.x0 + bx.x1) / 2 + e),
                         r(bx.y1 - bx.y0 if where == "ceiling" else bx.z1 - bx.z0)])
    wcsv(out / "data" / "battens.csv", ["id", "wall", "measured", "centre_mm", "length_mm"], rows)


def write(m, out: Path) -> None:
    layouts = nest(all_parts(m))
    m["layouts"] = layouts
    write_tables(m, out, layouts)
    write_dims(m, out)
    write_dxf(out, layouts)
    write_model(m, out)
    write_drawings(m, out, layouts)
